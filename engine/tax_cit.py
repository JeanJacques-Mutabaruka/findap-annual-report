"""Taxable income (detail + Excel with live formulas), RRA CIT annex view (BS / P&L with the RRA totals)
and the code <-> RRA mapping table.

Nothing here changes a figure: the taxable income is the one computed by build_model.income_tax (the rules
of the Excel tool V11); this module explains it line by line and rebuilds it with Excel formulas.
"""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd

from engine import rra as rra_mod
from engine.code_suggest import normal_side

MONEY = '#,##0;[Red](#,##0);"-"'
PCT = "0.00%"
COMPUTED = {"BS 05.08": "net_profit", "BS 08.02.01": "tax_charge"}

# Why each amount is added back. The coefficients and lines come from the Excel tool V11 (Incometax sheet);
# the legal references are to be completed by the firm — the app does not invent them.
ADD_BACK_RULES = {
    "PL 04.05": "Bad debts written off: added back in full (V11 rule).",
    "PL 04.08": "Communication: 20% added back (V11 rule — private-use share).",
    "PL 04.09": "Entertainment: added back in full (V11 rule).",
    "PL 04.16": "Other utilities: 20% added back (V11 rule — private-use share).",
    "PL 04.17": "Fuel: 20% added back (V11 rule — private-use share).",
    "PL 05.08.02": "Fines and penalties: added back in full (V11 rule).",
    "PL 05.08.03": "Personal consumption: added back in full (V11 rule).",
    "PL 05.08.04": "Non supported expenses (no EBM / supporting document): added back in full (V11 rule).",
    "PL 05.08.05": "Management fees paid: only the part above the limit (% of turnover) is added back.",
    "PL 07.01.01": "Interest to related parties: added back in full (V11 rule).",
}


def _side_sign(code: str) -> int:
    """Presentation sign in the RRA annex: +1 = shown as debit − credit, −1 = shown as credit − debit."""
    return -1 if normal_side(code) == "CR" else 1


def chart_lines(chart: dict) -> list[dict]:
    out = []
    for part, stmt in (("bs", "Balance sheet"), ("pnl", "P&L")):
        for sec in chart[part]:
            for g in sec.get("groups", []):
                for ln in g["lines"]:
                    out.append({**ln, "statement": stmt, "section": sec.get("title"), "group": g.get("label") or sec.get("title"),
                                "group_id": g.get("id"), "sign": g.get("sign")})
    return out


# ============================================================ mapping table ===
def mapping_df(chart: dict, rra: dict) -> pd.DataFrame:
    by_key = {rra_mod.key(x): x for x in rra["lines"]}
    rows = []
    for ln in chart_lines(chart):
        r = ln.get("rra")
        hit = by_key.get(f"{r['sheet']}:{r['row']}") if r else None
        rows.append({"CIT code": ln["code"], "Statement line (report)": ln["label"], "Report group": ln["group"],
                     "RRA sheet": {"BS": "BALANCE_SHEET", "PL": "PROFIT_AND_LOSS"}.get(r["sheet"], "") if r else "",
                     "RRA row": r["row"] if r else None, "RRA Serial No": r["serial"] if r else "",
                     "RRA description": hit["desc"] if hit else "",
                     "Link": ("sub-code → rolls up into the RRA line" if r.get("rollup") else "same line") if r else "⚠️ NO RRA LINE — to decide",
                     "Old V11 code": ln.get("old_code") or "(new)", "Note": ln.get("rra_note", "")})
    return pd.DataFrame(rows)


# ====================================================== taxable income detail ===
def taxable_income_df(model: dict) -> pd.DataFrame:
    """One row per step of the taxable-income computation, CY and PY, with how each amount is obtained."""
    det = model["income_tax"]["detail"]
    ys = [y for y in ("cy", "py") if y in det]
    c = det["cy"]
    p = det.get("py")
    rows = []

    def add(step, cy, py=None, how="", code="", coef=None, kind="line"):
        rows.append({"Step": step, "Code": code, "Coefficient": coef, "CY": cy,
                     "PY": py if "py" in ys else None, "How it is computed": how, "kind": kind})

    g = lambda k: p[k] if p else None
    add("Profit before tax", c["pbt"], g("pbt"), "From the P&L: all P&L codes, credit − debit (revenues − expenses).")
    add("Grant received", c["grant"], g("grant"), "Kept at 0 for the moment — the treatment of grants is to be designed.",
        kind="input")
    add("Profit net of grant received", c["pn"], g("pn"), "Profit before tax + grant received.", kind="total")
    add("Non-admissible expenses (added back)", None, None, "", kind="group")
    for i, a in enumerate(c["add_detail"]):
        pa = p["add_detail"][i]["amount"] if p else None
        if a["kind"] == "mgmt_fee":
            add(f"   {a['label']}", a["amount"], pa,
                f"max(0, management fees {a['code']} {a['fees']:,.0f} − {a['limit_rate']:.0%} × turnover {a['turnover']:,.0f} "
                f"(all revenues: codes {' + '.join(x + '…' for x in a['turnover_prefixes'])}, RRA P&L 1 + 8) = {a['limit']:,.0f}) = {a['amount']:,.0f}.", code=a["code"])
        else:
            add(f"   {a['label']}", a["amount"], pa, f"{a['coef']:.0%} × balance of {a['code']} (debit − credit) = "
                f"{a['coef']:.0%} × {a['base']:,.0f}. " + ADD_BACK_RULES.get(a["code"], ""), code=a["code"], coef=a["coef"])
    add("Total expenses added back", c["tot_add"], g("tot_add"), "Sum of the lines above (RRA P&L 11).", kind="total")
    add("Adjusted profit", c["adj"], g("adj"), "Profit net of grant + total added back.", kind="total")
    add("Non-taxable income (deducted)", None, None, "", kind="group")
    for i, d in enumerate(c["deductions"]):
        pd_ = -p["deductions"][i]["amount"] if p else None
        cap = "the whole amount" if d["cap"] is None else f"up to {d['cap']:,.0f}"
        add(f"   {d['label']}", -d["amount"], pd_, f"Income of {d['code']} (credit − debit) {d['income']:,.0f}, deducted "
            f"{cap} (RRA P&L {d['rra_line']}).", code=d["code"])
    add("Taxable income before losses", c["before_losses"], g("before_losses"), "Adjusted profit − non-taxable income.",
        kind="total")
    add("Tax losses available (brought forward)", c["loss_available"], g("loss_available"),
        f"Losses of the previous years still usable — {c['lb_source']}.", kind="input")
    add("Tax losses deducted", c["lb"], g("lb"), "min(losses available, taxable income before losses if positive) — "
        "RRA P&L 14.", kind="line")
    add("Tax base", c["base"], g("base"), "Taxable income before losses − losses deducted.", kind="total")
    if c["company_type"] == "CORPORATE":
        add("CIT rate", c["rate"], g("rate"), "Entered on 3 · Company & Report Data (per-year tax data).", kind="rate")
        add("Tax charge", c["charge"], g("charge"), "Tax base × CIT rate if the base is positive, otherwise 0.", kind="total")
    else:
        add("Tax charge", c["charge"], g("charge"), "Individual brackets (chart_of_accounts.json → income_tax).", kind="total")
    add("Quarterly prepayments", c["prepay"], g("prepay"), "Entered on 3 · Company & Report Data.", kind="input")
    add("Withholding tax (3% & 15%)", c["wht"], g("wht"), "Entered on 3 · Company & Report Data.", kind="input")
    add("Income tax payable", c["payable"], g("payable"), "Tax charge − prepayments − withholding tax.", kind="grandtotal")
    return pd.DataFrame(rows)


def tax_warnings(model: dict) -> list[str]:
    det = model["income_tax"]["detail"]
    w = []
    for y, lab in (("cy", "CY"), ("py", "PY")):
        if y not in det:
            continue
        d = det[y]
        if "no tax losses entered" in d["lb_source"] and d["before_losses"] < 0:
            w.append(f"{lab}: tax loss of the year and no losses of previous years entered — enter the tax losses per "
                     "year of origin (3 · Company & Report Data → Period & tax).")
        if "older project" in d["lb_source"]:
            w.append(f"{lab}: losses entered as a total per year (older project) — enter them per year of origin so that "
                     "the 5-year limit is applied.")
        if d["lb"] < 0:
            w.append(f"{lab}: {-d['lb']:,.0f} of tax losses deducted — RRA P&L line 15 (10 + 11 − 12 − 13) does not deduct "
                     "them: RRA 15 is higher than the tax base by this amount; income tax (RRA 16) is computed after losses.")
        if y == "cy" and d["prepay"] == 0 and d["charge"] > 0:
            w.append("CY: no quarterly prepayments entered while a tax charge exists — check the prepayments paid.")
    w.append("Depreciation: the book depreciation of the accounts is not replaced by the tax depreciation (RRA reducing "
             "balance per group of assets). Differences will be computed with the fixed-asset register.")
    return w


# ============================================================== RRA annex ===
def rra_values(chart: dict, rra: dict, tb_rows: list[dict], model: dict) -> dict:
    """{"cy": {key: value}, "py": {...}} values of every RRA line for the two years of a model."""
    lines = {ln["code"]: ln for ln in chart_lines(chart)}
    det = model["income_tax"]["detail"]
    kf = model["key_figures"]
    out = {}
    for y in ("cy", "py"):
        if y not in det:
            continue
        leaf: dict[str, float] = {}
        for r in tb_rows:
            ln = lines.get(r.get("code"))
            if not ln or not ln.get("rra") or ln["code"] in COMPUTED:
                continue
            k = f"{ln['rra']['sheet']}:{ln['rra']['row']}"
            net = float(r.get(f"debit_{y}") or 0) - float(r.get(f"credit_{y}") or 0)
            leaf[k] = leaf.get(k, 0.0) + net * _side_sign(ln["code"])
        for code, what in COMPUTED.items():
            ln = lines.get(code)
            if ln and ln.get("rra"):
                v = kf["net_profit"][y] if what == "net_profit" else model["income_tax"]["charge"][y]
                leaf[f"{ln['rra']['sheet']}:{ln['rra']['row']}"] = v
        tax_rows = {x["serial"]: rra_mod.key(x) for x in rra["lines"] if x["sheet"] == "PL" and x["kind"] == "tax"}
        extra = {}
        if "11" in tax_rows:
            extra[tax_rows["11"]] = det[y]["tot_add"]
        for d in det[y]["deductions"]:
            if d["rra_line"] in tax_rows:
                extra[tax_rows[d["rra_line"]]] = extra.get(tax_rows[d["rra_line"]], 0.0) + d["amount"]
        if "14" in tax_rows:
            extra[tax_rows["14"]] = -det[y]["lb"]
        if "16" in tax_rows:
            extra[tax_rows["16"]] = det[y]["charge"]
        out[y] = rra_mod.evaluate(rra, leaf, extra)
    return out


def rra_df(chart: dict, rra: dict, tb_rows: list[dict], model: dict) -> pd.DataFrame:
    vals = rra_values(chart, rra, tb_rows, model)
    feeders: dict[str, list[str]] = {}
    for ln in chart_lines(chart):
        if ln.get("rra"):
            feeders.setdefault(f"{ln['rra']['sheet']}:{ln['rra']['row']}", []).append(ln["code"])
    subs = sub_code_values(chart, tb_rows, model)
    rows = []
    for x in rra["lines"]:
        k = rra_mod.key(x)
        rows.append({"Sheet": x["sheet"], "RRA row": x["row"], "Serial No": x["serial"], "Description": x["desc"],
                     "Kind": {"line": "line", "total": "total", "tax": "tax computation"}[x["kind"]],
                     "CY": vals["cy"].get(k), "PY": vals.get("py", {}).get(k) if "py" in vals else None,
                     "Formula": x.get("formula", ""), "Formula source": x.get("formula_source", ""),
                     "App codes": ", ".join(feeders.get(k, []))})
        for code, lab, vcy, vpy in subs.get(k, []):
            rows.append({"Sheet": x["sheet"], "RRA row": x["row"], "Serial No": "", "Description": f"   ↳ of which: {lab} ({code})",
                         "Kind": "of which", "CY": vcy, "PY": vpy if "py" in vals else None, "Formula": "",
                         "Formula source": "detail — included in the line above", "App codes": code})
    return pd.DataFrame(rows)


def sub_code_values(chart: dict, tb_rows: list[dict], model: dict) -> dict[str, list]:
    """{rra key: [(code, label, CY, PY)]} for the sub-codes (rollup) — shown as 'of which' under their RRA line."""
    out: dict[str, list] = {}
    single = model["meta"].get("single_year")
    for ln in chart_lines(chart):
        r = ln.get("rra")
        if not r or not r.get("rollup"):
            continue
        v = {}
        for y in ("cy", "py"):
            v[y] = sum((float(t.get(f"debit_{y}") or 0) - float(t.get(f"credit_{y}") or 0)) * _side_sign(ln["code"])
                       for t in tb_rows if t.get("code") == ln["code"])
        if abs(v["cy"]) > 0.5 or (not single and abs(v["py"]) > 0.5):
            out.setdefault(f"{r['sheet']}:{r['row']}", []).append((ln["code"], ln["label"], v["cy"], v["py"]))
    return out


def reconciliation_df(rra_table: pd.DataFrame, model: dict) -> pd.DataFrame:
    """Report figures vs the same figures rebuilt with the RRA structure."""
    def rv(sheet, serial, col):
        hit = rra_table[(rra_table["Sheet"] == sheet) & (rra_table["Serial No"] == serial)]
        return float(hit.iloc[0][col] or 0) if len(hit) else None
    kf, tax = model["key_figures"], model["income_tax"]
    items = [("Total assets", "BS", "4", kf["total_assets"]), ("Gross profit", "PL", "3", kf["gross_profit"]),
             ("Profit before tax", "PL", "10", kf["pbt"]), ("Income tax of the year", "PL", "16", tax["charge"])]
    rows = []
    for lab, sh, ser, rep in items:
        for y, yl in (("cy", "CY"), ("py", "PY")):
            if rep.get(y) is None or (y == "py" and model["meta"].get("single_year")):
                continue
            r = rv(sh, ser, yl)
            rows.append({"Figure": lab, "Year": yl, "Annual report": rep[y], "RRA line": f"{sh} {ser}", "RRA annex": r,
                         "Difference": None if r is None else r - rep[y]})
    bal = rv("BS", "10", "CY")
    rows.append({"Figure": "Total assets − (equity + liabilities) — must be 0", "Year": "CY",
                 "Annual report": 0.0, "RRA line": "BS 10", "RRA annex": bal, "Difference": bal})
    return pd.DataFrame(rows)


# ============================================================== Excel files ===
def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    return {"h": Font(name="Arial", bold=True, color="FFFFFF", size=10), "hf": PatternFill("solid", fgColor="1F4E79"),
            "b": Font(name="Arial", bold=True, size=10), "n": Font(name="Arial", size=10),
            "inp": Font(name="Arial", size=10, color="0000FF"), "inpf": PatternFill("solid", fgColor="FFF2CC"),
            "tot": PatternFill("solid", fgColor="DDEBF7"), "gt": PatternFill("solid", fgColor="BDD7EE"),
            "wrap": Alignment(wrap_text=True, vertical="top"), "bd": Border(top=Side(style="thin", color="7F7F7F")),
            "i": Font(name="Arial", italic=True, size=9, color="595959")}


def _tb_sheet(wb, chart: dict, tb_rows: list[dict], years: tuple[int, int | None]):
    S = _styles()
    ws = wb.create_sheet("TB")
    lines = {ln["code"]: ln for ln in chart_lines(chart)}
    y_cy, y_py = years
    hdr = ["CIT code", "Account", "RRA key", "RRA Serial No", f"Debit {y_cy}", f"Credit {y_cy}", f"Debit {y_py}",
           f"Credit {y_py}", f"Net {y_cy} (Dr − Cr)", f"Net {y_py} (Dr − Cr)", "Sign in RRA annex"]
    ws.append(hdr)
    for c in ws[1]:
        c.font, c.fill, c.alignment = S["h"], S["hf"], S["wrap"]
    for i, r in enumerate(tb_rows, start=2):
        ln = lines.get(r.get("code")) or {}
        rr = ln.get("rra") or {}
        ws.append([r.get("code"), r.get("account"), f"{rr['sheet']}:{rr['row']}" if rr else "", rr.get("serial", ""),
                   float(r.get("debit_cy") or 0), float(r.get("credit_cy") or 0), float(r.get("debit_py") or 0),
                   float(r.get("credit_py") or 0), f"=E{i}-F{i}", f"=G{i}-H{i}", _side_sign(r.get("code") or "")])
        for c in ws[i]:
            c.font = S["n"]
        for col in "EFGHIJ":
            ws[f"{col}{i}"].number_format = MONEY
    n = len(tb_rows) + 1
    ws.append(["TOTAL", "", "", "", f"=SUM(E2:E{n})", f"=SUM(F2:F{n})", f"=SUM(G2:G{n})", f"=SUM(H2:H{n})",
               f"=SUM(I2:I{n})", f"=SUM(J2:J{n})"])
    for c in ws[n + 1]:
        c.font, c.fill, c.border = S["b"], S["tot"], S["bd"]
        c.number_format = MONEY
    for col, w in zip("ABCDEFGHIJK", (14, 40, 10, 12, 15, 15, 15, 15, 16, 16, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:K{n}"
    return ws


def tax_workbook(model: dict, tb_rows: list[dict], chart: dict, company: str = "", losses: dict | None = None) -> bytes:
    """Taxable income with live Excel formulas on the TB sheet (SUMIF by CIT code). Blue cells on yellow = inputs."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    S = _styles()
    det = model["income_tax"]["detail"]
    y_cy, y_py = model["meta"]["year_cy"], model["meta"]["year_py"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Taxable income"
    _tb_sheet(wb, chart, tb_rows, (y_cy, y_py))
    ws.append([f"{company} — Taxable income {y_cy} (CY) and {y_py} (PY)"])
    ws["A1"].font = Font(name="Arial", bold=True, size=13, color="1F4E79")
    ws.append([f"Generated {datetime.now():%d-%m-%Y %H:%M} — formulas read the TB sheet; blue figures on yellow are "
               f"inputs from 3 · Company & Report Data; columns G–H show the figures of the app for control."])
    ws["A2"].font = S["i"]
    ws.append([])
    ws.append(["Step", "Code", "Coefficient", f"CY {y_cy}", f"PY {y_py}", "How it is computed / rule",
               f"App CY {y_cy}", f"App PY {y_py}", "Check CY", "Check PY"])
    for c in ws[4]:
        c.font, c.fill, c.alignment = S["h"], S["hf"], S["wrap"]
    r = 5
    ref = {}

    def row(step, cy, py, how, code="", coef=None, kind="line", app=(None, None), key=None):
        nonlocal r
        ws.cell(r, 1, step); ws.cell(r, 2, code or None); ws.cell(r, 3, coef)
        ws.cell(r, 4, cy); ws.cell(r, 5, py); ws.cell(r, 6, how)
        ws.cell(r, 7, app[0]); ws.cell(r, 8, app[1])
        if app[0] is not None:
            ws.cell(r, 9, f'=IF(ABS(D{r}-G{r})<=1,"✓","≠")')
            ws.cell(r, 10, f'=IF(ABS(E{r}-H{r})<=1,"✓","≠")')
        for c in range(1, 11):
            cell = ws.cell(r, c)
            cell.font, cell.alignment = S["n"], S["wrap"]
            if c in (4, 5, 7, 8):
                cell.number_format = PCT if kind == "rate" else MONEY
        if coef is not None:
            ws.cell(r, 3).number_format = "0%"
            ws.cell(r, 3).font, ws.cell(r, 3).fill = S["inp"], S["inpf"]
        if kind in ("input", "rate"):
            for c in (4, 5):
                ws.cell(r, c).font, ws.cell(r, c).fill = S["inp"], S["inpf"]
        if kind in ("total", "grandtotal", "group"):
            for c in range(1, 7):
                ws.cell(r, c).font = S["b"]
                if kind != "group":
                    ws.cell(r, c).fill = S["gt"] if kind == "grandtotal" else S["tot"]
        if key:
            ref[key] = r
        r += 1

    c, p = det["cy"], det.get("py")
    P = lambda k: p[k] if p else None
    it = chart["income_tax"]
    row("Profit before tax", '=-SUMIF(TB!$A:$A,"PL *",TB!$I:$I)', '=-SUMIF(TB!$A:$A,"PL *",TB!$J:$J)',
        "Minus the SUM of (debit − credit) of every P&L code (PL …) in the TB sheet = revenues − expenses.",
        kind="total", app=(c["pbt"], P("pbt")), key="pbt")
    row("Grant received", 0, 0, "Kept at 0 — treatment of grants to be designed.", kind="input", key="grant")
    row("Profit net of grant received", f"=D{ref['pbt']}+D{ref['grant']}", f"=E{ref['pbt']}+E{ref['grant']}",
        "Profit before tax + grant.", kind="total", app=(c["pn"], P("pn")), key="pn")
    row("Non-admissible expenses (added back)", None, None, "", kind="group")
    first = r
    for i, a in enumerate(c["add_detail"]):
        pa = p["add_detail"][i]["amount"] if p else None
        if a["kind"] == "mgmt_fee":
            pres = a["turnover_prefixes"]
            tv = lambda col: "(" + "+".join(f'-SUMIF(TB!$A:$A,"{x}*",TB!${col}:${col})' for x in pres) + ")"
            row(a["label"], f'=MAX(0,SUMIF(TB!$A:$A,$B{r},TB!$I:$I)-$C{r}*{tv("I")})',
                f'=MAX(0,SUMIF(TB!$A:$A,$B{r},TB!$J:$J)-$C{r}*{tv("J")})',
                f"max(0, management fees − {a['limit_rate']:.0%} (column C) × turnover = credit − debit of all revenue "
                f"codes {', '.join(x + '…' for x in pres)} (RRA P&L 1 + 8)).", code=a["code"], coef=a["limit_rate"],
                app=(a["amount"], pa))
        else:
            row(a["label"], f'=SUMIF(TB!$A:$A,$B{r},TB!$I:$I)*$C{r}', f'=SUMIF(TB!$A:$A,$B{r},TB!$J:$J)*$C{r}',
                f"Coefficient × (debit − credit) of {a['code']}. " + ADD_BACK_RULES.get(a["code"], ""), code=a["code"],
                coef=a["coef"], app=(a["amount"], pa))
    last = r - 1
    row("Total expenses added back (RRA 11)", f"=SUM(D{first}:D{last})", f"=SUM(E{first}:E{last})", "Sum of the add-backs.",
        kind="total", app=(c["tot_add"], P("tot_add")), key="add")
    row("Adjusted profit", f"=D{ref['pn']}+D{ref['add']}", f"=E{ref['pn']}+E{ref['add']}",
        "Profit net of grant + total added back.", kind="total", app=(c["adj"], P("adj")), key="adj")
    row("Non-taxable income (deducted)", None, None, "", kind="group")
    f1 = r
    for i, d in enumerate(c["deductions"]):
        cap = d["cap"]
        fd = lambda col: (f'=-MAX(0,-SUMIF(TB!$A:$A,$B{r},TB!${col}:${col}))' if cap is None else
                          f'=-MIN({cap},MAX(0,-SUMIF(TB!$A:$A,$B{r},TB!${col}:${col})))')
        row(f"{d['label']} (RRA {d['rra_line']})", fd("I"), fd("J"),
            ("Income (credit − debit), deducted in full." if cap is None else f"Income (credit − debit), deducted up to {cap:,.0f}."),
            code=d["code"], app=(-d["amount"], -p["deductions"][i]["amount"] if p else None))
    l1 = r - 1
    row("Taxable income before losses", f"=D{ref['adj']}+SUM(D{f1}:D{l1})" if l1 >= f1 else f"=D{ref['adj']}",
        f"=E{ref['adj']}+SUM(E{f1}:E{l1})" if l1 >= f1 else f"=E{ref['adj']}", "Adjusted profit − non-taxable income.",
        kind="total", app=(c["before_losses"], P("before_losses")), key="bl")
    row("Tax losses available (brought forward)", c["loss_available"], P("loss_available"),
        f"Input — {c['lb_source']}. See the sheet 'Tax losses'.", kind="input", key="la")
    row("Tax losses deducted (RRA 14)", f"=-MIN(D{ref['la']},MAX(0,D{ref['bl']}))", f"=-MIN(E{ref['la']},MAX(0,E{ref['bl']}))",
        "min(losses available, taxable income before losses if positive).", kind="line", app=(c["lb"], P("lb")), key="lb")
    row("Tax base", f"=D{ref['bl']}+D{ref['lb']}", f"=E{ref['bl']}+E{ref['lb']}", "Taxable income before losses − losses deducted.",
        kind="total", app=(c["base"], P("base")), key="base")
    if c["company_type"] == "CORPORATE":
        row("CIT rate", c["rate"], P("rate"), "Input (per-year tax data).", kind="rate", key="rate")
        row("Tax charge (RRA 16)", f"=IF(D{ref['base']}>0,D{ref['base']}*D{ref['rate']},0)",
            f"=IF(E{ref['base']}>0,E{ref['base']}*E{ref['rate']},0)", "Tax base × rate if positive, otherwise 0.",
            kind="total", app=(c["charge"], P("charge")), key="charge")
    else:
        def prog(col):
            parts, lower = [], 0
            for upper, rate in it["individual_brackets_annual"]:
                top = f"{col}{ref['base']}" if upper is None else f"MIN({col}{ref['base']},{upper})"
                if rate:
                    parts.append(f"MAX(0,{top}-{lower})*{rate}")
                if upper is None:
                    break
                lower = upper
            return "=" + ("+".join(parts) or "0")
        row("Tax charge (RRA 16)", prog("D"), prog("E"), "Individual brackets from chart_of_accounts.json → income_tax.",
            kind="total", app=(c["charge"], P("charge")), key="charge")
    row("Quarterly prepayments", c["prepay"], P("prepay"), "Input.", kind="input", key="pre")
    row("Withholding tax (3% & 15%)", c["wht"], P("wht"), "Input.", kind="input", key="wht")
    row("Income tax payable", f"=D{ref['charge']}-D{ref['pre']}-D{ref['wht']}", f"=E{ref['charge']}-E{ref['pre']}-E{ref['wht']}",
        "Tax charge − prepayments − withholding tax.", kind="grandtotal", app=(c["payable"], P("payable")))
    if losses:
        wl = wb.create_sheet("Tax losses")
        wl.append(["Year of origin", "Tax loss", "Used before the TB years", "Source", "Usable until (year)"] +
                  [f"Used in {y}" for y in losses["years_list"]] + ["Remaining"])
        for cc in wl[1]:
            cc.font, cc.fill, cc.alignment = S["h"], S["hf"], S["wrap"]
        for o in losses["origins"]:
            wl.append([o["year"], o["loss"], o["used_before"], o["source"], o["expires_after"]] +
                      [o["used_by_year"].get(y, 0.0) for y in losses["years_list"]] + [o["remaining"]])
            for cc in wl[wl.max_row][1:]:
                cc.number_format = MONEY
            wl.cell(wl.max_row, 5).number_format = "0"
        wl.append([])
        wl.append([f"Carry-forward period: {losses['years']} years" + (" (extension for this client)" if losses.get("extension") else "")])
        for col, w in zip("ABCDEFGHIJKLMN", (14, 15, 18, 22, 14) + (13,) * 9):
            wl.column_dimensions[col].width = w
    for col, w in zip("ABCDEFGHIJ", (38, 13, 11, 16, 16, 70, 15, 15, 9, 9)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B5"
    rules = wb.create_sheet("Add-back rules")
    rules.append(["CIT code", "Line", "Coefficient / limit", "Rule used by the app", "Legal reference (to complete)"])
    for cc in rules[1]:
        cc.font, cc.fill = S["h"], S["hf"]
    for a in c["add_detail"]:
        rules.append([a["code"], a["label"], a.get("coef", a.get("limit_rate")), ADD_BACK_RULES.get(a["code"], ""), ""])
    for d in c["deductions"]:
        rules.append([d["code"], d["label"], None, f"Non-taxable income deducted (RRA P&L {d['rra_line']})" +
                      ("" if d["cap"] is None else f", up to {d['cap']:,.0f}"), ""])
    rules.append([])
    rules.append(["All rates, limits and codes come from data/chart_of_accounts.json → income_tax (one place)."])
    for rw in rules.iter_rows(min_row=2, min_col=3, max_col=3):
        rw[0].number_format = "0%"
    for col, w in zip("ABCDE", (13, 32, 11, 90, 40)):
        rules.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rra_workbook(chart: dict, rra: dict, tb_rows: list[dict], model: dict, company: str = "") -> bytes:
    """RRA BS and P&L annex with live formulas: lines = SUMIF on the TB sheet by RRA key × presentation sign;
    totals = the RRA formula written with cell references."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    S = _styles()
    y_cy, y_py = model["meta"]["year_cy"], model["meta"]["year_py"]
    table = rra_df(chart, rra, tb_rows, model)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, title in (("BS", "RRA Balance sheet"), ("PL", "RRA P&L")):
        ws = wb.create_sheet(title)
        ws.append([f"{company} — {title} (RRA CIT annex layout) — {y_cy} and {y_py}"])
        ws["A1"].font = Font(name="Arial", bold=True, size=13, color="1F4E79")
        ws.append(["Lines: SUMIF of the TB sheet on the RRA key × presentation sign · totals: RRA formula · values in "
                   "blue: from the app (profit of the year, tax computation)."])
        ws["A2"].font = S["i"]
        ws.append(["Serial No", "Description", f"CY {y_cy}", f"PY {y_py}", "Formula", "Formula source", "App codes"])
        for cc in ws[3]:
            cc.font, cc.fill, cc.alignment = S["h"], S["hf"], S["wrap"]
        sub = table[(table["Sheet"] == sheet) & (table["Kind"] != "of which")].reset_index(drop=True)
        lines = [x for x in rra["lines"] if x["sheet"] == sheet]
        subs = sub_code_values(chart, tb_rows, model)
        display = []                                   # ("rra", line, table row) / ("sub", code, label)
        for i, x in enumerate(lines):
            display.append(("rra", x, sub.iloc[i]))
            for code, lab, _, _ in subs.get(rra_mod.key(x), []):
                display.append(("sub", code, lab))
        start = 4
        cell_of = {rra_mod.key(d[1]): start + j for j, d in enumerate(display) if d[0] == "rra"}
        computed_keys = {}
        for ln in chart_lines(chart):
            if ln["code"] in COMPUTED and ln.get("rra"):
                computed_keys[f"{ln['rra']['sheet']}:{ln['rra']['row']}"] = ln["code"]
        for j, d in enumerate(display):
            rr = start + j
            if d[0] == "sub":
                code, lab = d[1], d[2]
                ws.cell(rr, 2, f"   ↳ of which: {lab} ({code}) — detail, included in the line above")
                sign = _side_sign(code)
                ws.cell(rr, 3, f'=SUMIF(TB!$A:$A,"{code}",TB!$I:$I)*{sign}')
                ws.cell(rr, 4, f'=SUMIF(TB!$A:$A,"{code}",TB!$J:$J)*{sign}')
                ws.cell(rr, 7, code)
                for col in range(1, 8):
                    ws.cell(rr, col).font = S["i"]
                    ws.cell(rr, col).alignment = S["wrap"]
                for col in (3, 4):
                    ws.cell(rr, col).number_format = MONEY
                continue
            x, srow = d[1], d[2]
            k = rra_mod.key(x)
            ws.cell(rr, 1, x["serial"]); ws.cell(rr, 2, x["desc"])
            ws.cell(rr, 5, x.get("formula") or ""); ws.cell(rr, 6, x.get("formula_source") or ""); ws.cell(rr, 7, srow["App codes"])
            if x["kind"] == "line" and k not in computed_keys:
                for col, net in ((3, "I"), (4, "J")):
                    ws.cell(rr, col, f'=SUMPRODUCT((TB!$C$2:$C$5000="{k}")*TB!${net}$2:${net}$5000*TB!$K$2:$K$5000)')
            elif x["kind"] == "total" and x.get("formula"):
                for col, L in ((3, "C"), (4, "D")):
                    parts = []
                    for sign, term in rra_mod._terms(x["formula"]):
                        if "~" in term:
                            a_, b_ = term.split("~")
                            ra = min(y["row"] for y in lines if y["serial"] == a_)
                            rb = max(y["row"] for y in lines if y["serial"] == b_)
                            lvl = a_.count(".") + 1
                            refs = [f"{L}{cell_of[rra_mod.key(y)]}" for y in lines if ra <= y["row"] <= rb and y["level"] == lvl]
                        else:
                            refs = [f"{L}{cell_of[rra_mod.key(y)]}" for y in lines if y["serial"] == term]
                        if refs:
                            parts.append(("-" if sign < 0 else "+") + ("(" + "+".join(refs) + ")" if len(refs) > 1 else refs[0]))
                    ws.cell(rr, col, "=" + ("".join(parts).lstrip("+") or "0"))
            else:                                   # computed by the app / tax computation / not computed
                ws.cell(rr, 3, srow["CY"]); ws.cell(rr, 4, srow["PY"])
            for col in range(1, 8):
                cell = ws.cell(rr, col)
                cell.font = S["b"] if x["kind"] == "total" else S["n"]
                cell.alignment = S["wrap"]
                if col in (3, 4):
                    cell.number_format = MONEY
            if x["kind"] == "total":
                for col in range(1, 5):
                    ws.cell(rr, col).fill = S["tot"]
            if k in computed_keys or x["kind"] == "tax":
                for col in (3, 4):
                    ws.cell(rr, col).font, ws.cell(rr, col).fill = S["inp"], S["inpf"]
        for col, w in zip("ABCDEFG", (10, 58, 17, 17, 34, 30, 30)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "C4"
    _tb_sheet(wb, chart, tb_rows, (y_cy, y_py))
    rec = reconciliation_df(table, model)
    ws = wb.create_sheet("Report vs RRA")
    ws.append(list(rec.columns))
    for cc in ws[1]:
        cc.font, cc.fill = S["h"], S["hf"]
    for rw in rec.itertuples(index=False):
        ws.append(list(rw))
    for rw in ws.iter_rows(min_row=2):
        for cc in (rw[2], rw[4], rw[5]):
            cc.number_format = MONEY
    for col, w in zip("ABCDEF", (48, 7, 18, 10, 18, 16)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def mapping_workbook(chart: dict, rra: dict) -> bytes:
    from engine.package import workbook
    return workbook({"Code to RRA": mapping_df(chart, rra),
                     "RRA lines": pd.DataFrame([{k: x.get(k) for k in ("sheet", "row", "serial", "desc", "kind", "formula",
                                                                          "formula_source")} for x in rra["lines"]])})


def compare_workbook(res: dict) -> bytes:
    from engine.package import workbook
    def lines(items, k):
        return pd.DataFrame([{"Sheet": i[k]["sheet"], "Row": i[k]["row"], "Serial No": i[k]["serial"],
                              "Description": i[k]["desc"], "App codes": ", ".join(i.get("codes", []))} for i in items])
    return workbook({"Proposals": pd.DataFrame(res["proposals"]), "Added in RRA": lines(res["added"], "new"),
                     "Removed from RRA": lines(res["removed"], "old"),
                     "Wording changed": pd.DataFrame([{"Sheet": c["old"]["sheet"], "Serial No": c["old"]["serial"],
                                                        "Old wording": c["old"]["desc"], "New wording": c["new"]["desc"],
                                                        "App codes": ", ".join(c["codes"])} for c in res["changed"]]),
                     "Rows moved": pd.DataFrame([{"Sheet": m["old"]["sheet"], "Serial No": m["old"]["serial"],
                                                   "Old row": m["old"]["row"], "New row": m["new"]["row"]} for m in res["moved"]])})
