"""Regression tests for the engine. Run with:  pytest -q

Expected figures are those of the V11 Excel generator on its own test data (TEST COMPANY LTD, FY 2025).
"""
from __future__ import annotations

import copy
import io
import json
import sys
import zipfile
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from engine import CHART, DEMO, NOTES, TEMPLATE, build_model, code_suggest, multiyear, package, render_report, tb_io  # noqa: E402

CHART_D = json.loads(CHART.read_text(encoding="utf-8"))
NOTES_D = json.loads(NOTES.read_text(encoding="utf-8"))


@pytest.fixture
def demo():
    return json.loads(DEMO.read_text(encoding="utf-8"))


def build(d):
    return build_model.build(d, CHART_D, NOTES_D)


def ids(model, level):
    return {c["id"] for c in model["controls"] if c["level"] == level}


def test_figures_match_excel_v11(demo):
    k = build(demo)["key_figures"]
    assert round(k["revenue"]["cy"]) == 535_837_403
    assert round(k["pbt"]["cy"]) == 127_715_819
    assert round(k["net_profit"]["cy"]) == 89_851_899
    assert round(k["total_assets"]["cy"]) == 296_521_579
    assert round(k["total_equity"]["cy"]) == 233_349_876


def test_demo_controls(demo):
    m = build(demo)
    assert "ACCDEP_DEBIT" in ids(m, "BLOCKING")
    assert {"TB_PLUG_COMMENT", "CIT_RATE"} <= ids(m, "WARNING")
    assert {"TB_BALANCED", "BS_BALANCED", "CF_CASH"} <= ids(m, "INFO")


def test_unbalanced_tb_is_blocking(demo):
    d = copy.deepcopy(demo)
    d["trial_balance"][0]["debit_cy"] += 1000
    m = build(d)
    assert "TB_NOT_BALANCED" in ids(m, "BLOCKING")


def test_unmapped_code_is_blocking(demo):
    d = copy.deepcopy(demo)
    d["trial_balance"][0]["code"] = "XX 9.9"
    assert "TB_UNMAPPED" in ids(build(d), "BLOCKING")


VALID = {c["code"] for s in CHART_D["pnl"] + CHART_D["bs"] for g in s.get("groups", []) for c in g["lines"]}


def test_tb_balance_and_line_issues(demo):
    tb = tb_io.from_pair_rows(demo["trial_balance"], 2025)
    assert tb_io.years_of(tb) == [2024, 2025]
    assert tb_io.balance(tb)["ok"]
    iss = tb_io.line_issues(tb, VALID)
    assert (iss["Issue"].str.contains("Accumulated depreciation")).sum() == 8   # 4 lines x 2 years
    assert not iss["Issue"].str.contains("Identical line").any()


def test_upload_mapping_roundtrip():
    data = (ROOT / "data" / "demo" / "demo_trial_balance.xlsx").read_bytes()
    name = "demo_trial_balance.xlsx"
    hdr = tb_io.guess_header_row(data, name)
    assert hdr == 3
    raw = tb_io.read_table(data, name, 0, hdr)
    mp = tb_io.guess_mapping(list(raw.columns))
    assert mp["account"] == "Account name" and [y["year"] for y in mp["years"]] == [2025, 2024]
    assert mp["years"][1]["credit"] == "Credit 2024"
    tb, notes = tb_io.to_tb(raw, mp)
    assert len(tb) == 42 and tb_io.balance(tb)["ok"]
    assert (tb["code"] == "").sum() == 4


def test_five_year_upload_drops_unlabelled_total():
    data = (ROOT / "data" / "demo" / "demo_trial_balance_5years.xlsx").read_bytes()
    name = "x.xlsx"
    raw = tb_io.read_table(data, name, 0, tb_io.guess_header_row(data, name))
    mp = tb_io.guess_mapping(list(raw.columns))
    assert sorted(y["year"] for y in mp["years"]) == [2021, 2022, 2023, 2024, 2025]
    tb, notes = tb_io.to_tb(raw, mp)
    assert len(tb) == 42 and tb_io.balance(tb)["ok"]
    assert any("total" in n for n in notes)


def test_nan_account_does_not_crash():
    df = pd.DataFrame({"Account": ["Bank", None], "Debit 2025": [100, 5], "Credit 2025": [0, 5]})
    tb, _ = tb_io.to_tb(df, tb_io.guess_mapping(list(df.columns)))
    iss = tb_io.line_issues(tb, VALID)
    assert "Account name missing" in set(iss["Issue"])
    assert code_suggest.suggest(float("nan"))[0] is None and code_suggest.suggest(None)[0] is None


def test_signed_balance_column():
    df = pd.DataFrame({"Account": ["Bank", "Capital"], "Balance 2025": [100, -100], "Balance 2024": ["(50)", "50"]})
    tb, _ = tb_io.to_tb(df, tb_io.guess_mapping(list(df.columns)))
    assert tb.loc[0, "debit_2025"] == 100 and tb.loc[1, "credit_2025"] == 100 and tb.loc[0, "credit_2024"] == 50


def test_multiyear_pairs_match_two_year_engine(demo):
    proj, tb = multiyear.from_pair_input(demo)
    m = build(multiyear.pair_input(proj, tb, 2025))
    assert round(m["key_figures"]["net_profit"]["cy"]) == 89_851_899
    assert m["meta"]["cit_rate"] if "cit_rate" in m["meta"] else True
    five = json.loads((ROOT / "data" / "demo" / "demo_project_5years.json").read_text(encoding="utf-8"))
    tb5 = tb_io.from_records(five["trial_balance"])
    assert multiyear.pairs(tb_io.years_of(tb5)) == [2022, 2023, 2024, 2025]
    models = {y: build(multiyear.pair_input(five, tb5, y)) for y in multiyear.pairs(tb_io.years_of(tb5))}
    rows, years = multiyear.multi_statement(models, "pnl")
    assert years == [2021, 2022, 2023, 2024, 2025]
    npr = next(r for r in rows if r["label"].startswith("Net profit"))
    assert round(npr["values"][-1]) == 89_851_899
    notes, ny = multiyear.multi_notes(models)
    assert ny == years and notes[0]["rows"]


def test_tb_template_reuploads():
    cat = pd.DataFrame([{"code": "BS 01.01", "statement": "Balance sheet", "section": "x", "group": "g", "line": "Land", "note": "11"}])
    x = package.tb_template([2025, 2024, 2023], cat, "ACME")
    hdr = tb_io.guess_header_row(x, "t.xlsx", "Trial Balance")
    assert hdr == 3
    raw = tb_io.read_table(x, "t.xlsx", "Trial Balance", hdr)
    mp = tb_io.guess_mapping(list(raw.columns))
    assert [y["year"] for y in mp["years"]] == [2025, 2024, 2023] and mp["code"] == "CIT code"


def test_code_suggestions():
    assert code_suggest.suggest("ACCUMULATED DEPRECIATION")[0] == "BS 01.10"
    assert code_suggest.suggest("Bank charges")[0] == "PL 07.02"
    assert code_suggest.suggest("BPR BANK PLC RWF")[0] == "BS 03.01.03.02"
    assert code_suggest.suggest("Dep-Moto Vehicle")[0] == "PL 05.01"


def test_render_docx_no_placeholder_left(demo):
    m = build(demo)
    with pytest.raises(render_report.BlockingControls):
        render_report.render(str(TEMPLATE), m, io.BytesIO())
    out = io.BytesIO()
    rep = render_report.render(str(TEMPLATE), m, out, allow_blocking=True)
    assert not rep["missing"] and rep["draft"]
    xml = zipfile.ZipFile(io.BytesIO(out.getvalue())).read("word/document.xml").decode("utf-8")
    assert "{" not in "".join(__import__("re").findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))


def test_statements_workbook(demo):
    x = package.statements_workbook(build(demo))
    sheets = pd.ExcelFile(io.BytesIO(x)).sheet_names
    assert {"P&L", "Balance sheet", "Cash flow", "Notes", "Controls"} <= set(sheets)


def test_template_cascade_lists_and_code_derivation():
    import openpyxl
    rows = []
    for stt, key in (("P&L", "pnl"), ("Balance sheet", "bs")):
        for sec in CHART_D[key]:
            for g in sec.get("groups", []):
                for ln in g["lines"]:
                    rows.append({"code": ln["code"], "statement": stt, "section": sec.get("title"),
                                 "group": g.get("label") or sec.get("title"), "line": ln["label"], "note": ""})
    cat = pd.DataFrame(rows)
    x = package.tb_template([2025, 2024], cat, "ACME")
    wb = openpyxl.load_workbook(io.BytesIO(x))
    assert wb.sheetnames == ["Instructions", "Trial Balance", "CIT codes", "Lists"]
    assert wb["Lists"].sheet_state == "hidden"
    assert len(wb["Trial Balance"].data_validations.dataValidation) == 5
    ws = wb["Trial Balance"]
    ws["C30"], ws["D30"], ws["F30"], ws["G30"] = "Cash and Cash equivalents", "Bank balances", "BK USD", 100
    buf = io.BytesIO(); wb.save(buf); data = buf.getvalue()
    raw = tb_io.read_table(data, "t.xlsx", "Trial Balance", tb_io.guess_header_row(data, "t.xlsx", "Trial Balance"))
    mp = tb_io.guess_mapping(list(raw.columns))
    tb, notes = tb_io.to_tb(raw, mp, cat)
    assert tb.loc[tb["account"] == "BK USD", "code"].iloc[0] == "BS 03.01.03.02"


def test_duplicate_descriptions_blocking_and_merge(demo):
    tb = tb_io.from_pair_rows(demo["trial_balance"], 2025)
    groups = tb_io.duplicate_groups(tb)
    acc = [g for g in groups if g["key"] == "accumulated depreciation"]
    assert acc and len(acc[0]["rows"]) == 4 and acc[0]["same_code"]
    assert "DUPLICATE_ACCOUNT" in ids(build(demo), "BLOCKING")
    merged = tb_io.merge_lines(tb, acc[0]["rows"])
    assert len(merged) == len(tb) - 3
    assert abs(merged[tb_io.amount_cols(merged)].sum() - tb[tb_io.amount_cols(tb)].sum()).max() < 0.01
    assert tb_io.balance(merged)["ok"]
    assert tb_io.desc_key("Bank  of Kigali - RWF ") == tb_io.desc_key("BANK OF KIGALI RWF")


def test_template_registry_and_generic_bookmarks(demo):
    import re as _re
    from engine import templates
    ts = templates.list_templates()
    assert ts and ts[0]["id"] == "standard-v09" and not ts[0]["error"]
    assert templates.check(ts[0]["docx"])["ok"]
    # a template using the generic names: rename two V09 bookmarks
    z = zipfile.ZipFile(str(TEMPLATE))
    parts = {n: z.read(n) for n in z.namelist()}
    x = parts["word/document.xml"].decode("utf-8")
    x = x.replace('w:name="str_Coverpage_Companyname"', 'w:name="str_companyname__9"')
    x = x.replace('w:name="tbl_PnL_Pnltable"', 'w:name="tbl_pnl"').replace('w:name="var_Pnlrelatednotessection"', 'w:name="var_notes"')
    parts["word/document.xml"] = x.encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zo:
        for n, d in parts.items():
            zo.writestr(n, d)
    chk = templates.check(buf.getvalue())
    assert chk["ok"] and "companyname" in chk["fields"]
    out = io.BytesIO()
    rep = render_report.render(io.BytesIO(buf.getvalue()), build(demo), out, allow_blocking=True, font_name="Georgia")
    assert "str_companyname__9" in rep["replaced"] and "tbl_pnl" in rep["replaced"] and "var_notes" in rep["replaced"]
    doc = zipfile.ZipFile(io.BytesIO(out.getvalue())).read("word/document.xml").decode("utf-8")
    assert set(_re.findall(r'w:ascii="([^"]+)"', doc)) <= {"Georgia", "Symbol"}


def test_company_data_roundtrip_and_multi_tb(demo):
    from engine import company_data
    proj, tb = multiyear.from_pair_input(demo)
    proj["general"]["audit_company_name"] = "AUDIT X"
    proj.setdefault("years_data", {}).setdefault("2025", {})["prepayments"] = 1234.0
    for fmt, data in (("json", company_data.to_json(proj)), ("xlsx", company_data.to_excel(proj, [2024, 2025]))):
        part = company_data.from_file(data, f"x.{fmt}")
        assert part["general"]["audit_company_name"] == "AUDIT X"
        assert abs(part["years_data"]["2025"]["prepayments"] - 1234) < 0.01
        assert abs(part["years_data"]["2025"]["cit_rate"] - proj["years_data"]["2025"]["cit_rate"]) < 1e-6
        target = {"meta": {}, "general": {}, "years_data": {}}
        assert "printed data" in company_data.apply(target, part)
    cat = pd.DataFrame([{"code": "BS 03.01.03.02", "statement": "Balance sheet", "section": "s", "group": "g", "line": "Bank", "note": ""},
                        {"code": "BS 05.01", "statement": "Balance sheet", "section": "s", "group": "g2", "line": "Capital", "note": ""}])
    rows = [{"code": "BS 03.01.03.02", "account": "BANK", "amounts": {2025: (100, 0), 2024: (50, 0)}},
            {"code": "BS 05.01", "account": "CAPITAL", "amounts": {2025: (0, 100), 2024: (0, 50)}}]
    x = package.tb_template([2025, 2024], cat, "ACME", sheets={"TB A": rows, "TB B": rows[:1]})
    found = tb_io.tb_sheets(x, "t.xlsx")
    assert [f["sheet"] for f in found] == ["TB A", "TB B"]
    assert found[0]["balanced"] and not found[1]["balanced"] and found[0]["lines"] == 2


def test_single_year_and_detail_levels(demo):
    proj, tb = multiyear.from_pair_input(demo)
    tb1 = tb.drop(columns=["debit_2024", "credit_2024"])            # company in its first year
    assert multiyear.single_year(proj, tb1, 2025)
    m = build(multiyear.pair_input(proj, tb1, 2025))
    assert m["meta"]["single_year"] and len(m["equity"]["blocks"]) == 1
    assert "RE_ROLLFORWARD" not in ids(m, "WARNING")
    out = io.BytesIO()
    render_report.render(str(TEMPLATE), m, out, allow_blocking=True)
    xml = zipfile.ZipFile(io.BytesIO(out.getvalue())).read("word/document.xml").decode("utf-8")
    assert ">2025<" in xml and ">2024<" not in xml
    # forced comparative
    proj["meta"]["comparative"] = "with"
    assert not multiyear.single_year(proj, tb1, 2025)
    # levels
    m2 = build(demo)
    det = render_report.apply_level(m2["pnl"], "detailed")
    summ = render_report.apply_level(m2["pnl"], "summarised")
    cond = render_report.apply_level(m2["bs"], "condensed")
    assert len(det) > len(summ) > 0
    assert [r["label"] for r in cond if r["type"] == "line"] == ["Non Current Assets", "Current Assets", "Equity",
                                                                   "Non Current Liabilities", "Current Liabilities"]
    np_row = next(r for r in render_report.apply_level(m2["pnl"], "condensed") if r["type"] == "grandtotal")
    assert round(np_row["cy"]) == 89_851_899
    rep = render_report.render(str(TEMPLATE), {**m2, "meta": {**m2["meta"], "detail_level": "condensed"}}, io.BytesIO(),
                               allow_blocking=True)
    assert not rep["missing"]


def test_company_sheets_inside_tb_file(demo):
    from engine import company_data
    proj, tb = multiyear.from_pair_input(demo)
    proj["general"]["audit_company_name"] = "AUDIT Y"
    cat = pd.DataFrame([{"code": "BS 05.01", "statement": "Balance sheet", "section": "s", "group": "g", "line": "Capital", "note": ""}])
    x = package.tb_template([2025, 2024], cat, "ACME", extra_sheets=company_data.frames(proj, [2024, 2025])[0])
    names = tb_io.sheet_names(x, "t.xlsx")
    assert company_data.has_company_sheets(names)
    assert [f["sheet"] for f in tb_io.tb_sheets(x, "t.xlsx")] == ["Trial Balance"]
    assert company_data.from_file(x, "t.xlsx")["general"]["audit_company_name"] == "AUDIT Y"


# ------------------------------------------------------------------ V1-0j --
def _catalogue():
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    return pd.DataFrame([{"code": l["code"], "statement": "P&L" if k == "pnl" else "Balance sheet",
                          "section": s.get("title"), "group": g.get("label") or s.get("title"), "line": l["label"],
                          "note": g.get("note") or s.get("note") or ""}
                         for k in ("pnl", "bs") for s in ch[k] for g in s.get("groups", []) for l in g["lines"]])


def test_proposals_use_balance_direction():
    cat = _catalogue()
    p = code_suggest.propose
    assert p("BPR BANK PLC RWF", 3000, cat)["code"] == "BS 03.01.03.02" and p("BPR BANK PLC RWF", 3000, cat)["confidence"] == "High"
    od = p("BPR BANK PLC RWF", -3000, cat)
    assert od["code"] == "BS 08.01.04" and od["confidence"] == "Medium" and od["alternative"] == "BS 03.01.03.02" and od["comment"]
    assert p("Dep-Moto Vehicle", -200, cat)["code"] == "BS 01.10"
    assert p("Interest", -50, cat)["code"] == "PL 08.01" and p("Interest", 50, cat)["code"] == "PL 07.01"
    assert p("VAT", -99, cat)["code"] == "BS 08.01.06" and p("VAT", 99, cat)["code"] == "BS 03.01.04.01"
    assert p("Fines&Penality-PAYE", 50, cat)["code"] == "PL 05.08.02"
    assert p("Bank charges", 10, cat)["confidence"] == "High"            # 'bank' inside 'bank charges' is no doubt
    assert p("Office expenses", 10, cat)["confidence"] == "Medium"       # generic rule
    assert p("Misc", 10, cat)["code"] is None and p(None)["code"] is None
    assert code_suggest.normal_side("PL 08.01") == "CR" and code_suggest.normal_side("PL 02.05.03") == "CR"
    assert code_suggest.normal_side("BS 01.10") == "CR" and code_suggest.normal_side("BS 03.01.03.02") == "DR"


def test_zero_lines_and_delete():
    df = pd.DataFrame({"code": ["BS 03.01.03.02", "PL 04.12", ""], "account": ["Bank", "Rent", "Old account"],
                       "comment": "", "note": "", "debit_2025": [100, 0, None], "credit_2025": [0, 0, None],
                       "debit_2024": [50, 0, 0], "credit_2024": [0, 0, 0]})
    tb = tb_io.normalise(df)
    assert tb_io.zero_lines(tb) == [1, 2]
    out = tb_io.delete_lines(tb, [1, 2])
    assert list(out["account"]) == ["Bank"] and tb_io.line_net(tb, 0) == 100 and tb_io.line_net(tb, 1) is None


def test_ai_prompt_lists_every_code_and_output_reuploads(demo):
    from engine import ai_prompt
    cat = _catalogue()
    txt = ai_prompt.build(cat)
    assert all(f"| {c} |" in txt for c in cat["code"])
    assert "`CIT code`" in txt and "AI confidence" in txt and "Confidentiality" in txt
    # simulate the file returned by an AI: ledger code column, CIT code inserted before the name, 3 columns at the end
    rows = [{"Code": 1000 + i, "CIT code": r["code"], "Account name": r["account"],
             "Debit 2025": r["debit_cy"], "Credit 2025": r["credit_cy"], "Debit 2024": r["debit_py"],
             "Credit 2024": r["credit_py"], "AI statement line (check)": "x", "AI confidence": "High",
             "Comments": "Proposed … please check" if i == 0 else ""} for i, r in enumerate(demo["trial_balance"])]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as xw:
        pd.DataFrame(rows).to_excel(xw, index=False, startrow=2, sheet_name="TB")
    data = buf.getvalue()
    hdr = tb_io.guess_header_row(data, "tb_mapped.xlsx", "TB")
    raw = tb_io.read_table(data, "tb_mapped.xlsx", "TB", hdr)
    mp = tb_io.guess_mapping(list(raw.columns))
    assert mp["code"] == "CIT code" and mp["account"] == "Account name" and mp["comment"] == "Comments"
    tb, _ = tb_io.to_tb(raw, mp)
    assert len(tb) == len(rows) and tb_io.balance(tb)["ok"] and tb.at[0, "comment"].startswith("Proposed")
    assert list(tb["code"]) == [r["code"] for r in demo["trial_balance"]]


# ------------------------------------------------- 2026-09-27 RRA codes --
def test_old_codes_are_converted_on_upload():
    from engine import codes
    df = pd.DataFrame({"CIT code": ["BS 1.09", "BS 3.1.3.2", "PL 5.10", "PL 8.14", "BS 01.10"],
                       "Account": ["Acc dep", "Bank", "Fines", "Other income", "Acc dep new"],
                       "Debit 2025": [0, 100, 5, 0, 0], "Credit 2025": [50, 0, 0, 55, 0]})
    tb, notes = tb_io.to_tb(df, tb_io.guess_mapping(list(df.columns)))
    assert list(tb["code"]) == ["BS 01.10", "BS 03.01.03.02", "PL 05.08.02", "PL 08.13", "BS 01.10"]
    assert any("V11 numbering" in n and "BS 1.09 → BS 01.10" in n for n in notes)
    assert codes.to_new("BS 01.09") == "BS 01.09" and codes.from_serial("PL", "4.10") == "PL 04.10"
    assert all(codes.NEW_RE.match(c) for c in codes.OLD_TO_NEW.values())


def test_chart_links_every_code_to_the_rra_file():
    from engine import codes, rra
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    R = rra.load()
    keys = {rra.key(x): x for x in R["lines"]}
    lines = [l for k in ("pnl", "bs") for s in ch[k] for g in s.get("groups", []) for l in g["lines"]]
    assert len({l["code"] for l in lines}) == len(lines)
    for l in lines:
        assert codes.NEW_RE.match(l["code"]), l["code"]
        assert l["rra"], l["code"]                                            # every code is linked (28-09-2026)
        hit = keys[f"{l['rra']['sheet']}:{l['rra']['row']}"]
        assert hit["serial"] == l["rra"]["serial"] and hit["kind"] == "line"
        if l["rra"].get("rollup"):
            assert codes.to_serial(codes.parent(l["code"])) == l["rra"]["serial"]
    # every RRA line (not a total / tax line) has an app code
    linked = {f"{l['rra']['sheet']}:{l['rra']['row']}" for l in lines if l["rra"]}
    assert not [k for k, x in keys.items() if x["kind"] == "line" and k not in linked]


def test_rra_file_parsing_and_update_check():
    from engine import rra
    R = rra.load()
    pl = [x["serial"] for x in R["lines"] if x["sheet"] == "PL"]
    assert "4.10" in pl and "4.20" in pl and "8.10" in pl and pl.count("4.1") == 1
    assert [x["row"] for x in R["lines"] if x["sheet"] == "BS" and x["serial"] == "3.1.3.1"] == [32, 33]
    one = next(x for x in R["lines"] if x["sheet"] == "BS" and x["serial"] == "1")
    assert one["formula"] == "1.1 TO 1.9-1.10+1.11"
    new = copy.deepcopy(R)
    new["lines"] = [x for x in new["lines"] if not (x["sheet"] == "PL" and x["serial"] == "9.2")]
    next(x for x in new["lines"] if x["sheet"] == "BS" and x["serial"] == "1.7")["desc"] = "IT EQUIPMENT AND SOFTWARE"
    new["lines"].append({"sheet": "PL", "row": 200, "serial": "4.28", "raw": 4.28, "desc": "Subcontractors", "level": 2,
                         "kind": "line", "formula": "", "formula_source": ""})
    res = rra.compare(R, new, {"PL:99": ["PL 09.02"], "BS:9": ["BS 01.07"]})
    assert [a["new"]["serial"] for a in res["added"]] == ["4.28"]
    assert res["removed"][0]["codes"] == ["PL 09.02"] and res["changed"][0]["codes"] == ["BS 01.07"]
    acts = {p["action"] for p in res["proposals"]}
    assert {"ADD CODE", "RELINK CODE", "CHECK LABEL"} <= acts
    assert any(p["code"] == "PL 04.28" for p in res["proposals"])


def test_rra_annex_matches_the_report_and_tax_workbook(demo):
    from engine import rra, tax_cit
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    m = build(demo)
    table = tax_cit.rra_df(ch, rra.load(), demo["trial_balance"], m)
    rec = tax_cit.reconciliation_df(table, m)
    assert rec["Difference"].abs().max() < 1
    df = tax_cit.taxable_income_df(m)
    assert round(df.loc[df["Step"] == "Tax charge", "CY"].iloc[0]) == round(m["income_tax"]["charge"]["cy"])
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(tax_cit.tax_workbook(m, demo["trial_balance"], ch, "T")))
    f = [c.value for row in wb["Taxable income"].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("SUMIF(TB!" in x for x in f) and any(x.startswith("=IF(") for x in f)
    wb2 = load_workbook(io.BytesIO(tax_cit.rra_workbook(ch, rra.load(), demo["trial_balance"], m, "T")))
    assert {"RRA Balance sheet", "RRA P&L", "TB", "Report vs RRA"} <= set(wb2.sheetnames)


# ------------------------------------------------- 2026-09-28 tax corrections --
def test_management_fee_dividends_farming_and_losses(demo):
    d = copy.deepcopy(demo)
    d["trial_balance"] += [
        {"code": "PL 05.08.05", "account": "MGMT FEES", "debit_cy": 15_000_000, "credit_cy": 0, "debit_py": 0, "credit_py": 0},
        {"code": "PL 08.10", "account": "DIVIDENDS", "debit_cy": 0, "credit_cy": 2_000_000, "debit_py": 0, "credit_py": 0},
        {"code": "PL 08.02", "account": "FARMING", "debit_cy": 0, "credit_cy": 20_000_000, "debit_py": 0, "credit_py": 0}]
    d["tax"]["loss_available"] = {"cy": 5_000_000, "py": 0}
    m = build(d)
    c = m["income_tax"]["detail"]["cy"]
    turnover = sum(-(r["debit_cy"] - r["credit_cy"]) for r in d["trial_balance"] if r["code"].startswith(("PL 01", "PL 08")))
    assert abs(c["mgmt"]["amount"] - max(0, 15_000_000 - 0.02 * turnover)) < 1
    assert [x["amount"] for x in c["deductions"]] == [2_000_000, 12_000_000]          # dividends in full, farming capped
    assert c["lb"] == -5_000_000 and abs(c["base"] - (c["adj"] - 14_000_000 - 5_000_000)) < 1
    assert c["grant"] == 0
    assert any(x["id"] == "RRA_LINE15" for x in m["controls"])


def test_individual_brackets_from_one_table():
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    br = ch["income_tax"]["individual_brackets_annual"]
    assert build_model.progressive_tax(300_000, br) == 0
    assert abs(build_model.progressive_tax(1_000_000, br) - 0.2 * 640_000) < 1e-6
    assert abs(build_model.progressive_tax(2_000_000, br) - (0.2 * 840_000 + 0.3 * 800_000)) < 1e-6
    assert build_model.standard_cit_rate(ch, 2023) == 0.30 and build_model.standard_cit_rate(ch, 2025) == 0.28


def test_loss_schedule_five_years_fifo_and_extension():
    from engine import multiyear
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    # TB 2024-2025 with a profit of 10,000,000 each year (sales only)
    tb = tb_io.normalise(pd.DataFrame([{"code": "PL 01.01", "account": "Sales", "debit_2025": 0, "credit_2025": 10e6,
                                        "debit_2024": 0, "credit_2024": 10e6},
                                       {"code": "BS 03.01.03.02", "account": "Bank", "debit_2025": 10e6, "credit_2025": 0,
                                        "debit_2024": 10e6, "credit_2024": 0}]))
    proj = {"tax_losses": {"extension": False, "losses": [
        {"year": 2018, "loss": 3e6, "used_before": 0},        # usable until 2023 -> expired for 2024
        {"year": 2019, "loss": 4e6, "used_before": 1e6},      # 3e6 left, usable until 2024
        {"year": 2022, "loss": 9e6, "used_before": 0}]}}      # usable until 2027
    s = multiyear.loss_schedule(proj, tb, ch)
    assert s["by_year"][2024]["available"] == 12e6 and s["by_year"][2024]["used"] == 10e6
    o = {x["year"]: x for x in s["origins"]}
    assert o[2019]["used_by_year"][2024] == 3e6 and o[2022]["used_by_year"][2024] == 7e6   # oldest first
    assert s["by_year"][2025]["available"] == 2e6 and s["by_year"][2025]["used"] == 2e6
    proj["tax_losses"].update(extension=True, carry_forward_years=7)
    s7 = multiyear.loss_schedule(proj, tb, ch)
    assert s7["by_year"][2024]["available"] == 15e6                                        # 2018 still usable


def test_negative_amounts_and_balance_notices():
    tb = tb_io.normalise(pd.DataFrame([{"code": "PL 01.01", "account": "Sales", "debit_2025": -100, "credit_2025": 0},
                                       {"code": "BS 03.01.03.02", "account": "Bank", "debit_2025": 50, "credit_2025": 0}]))
    notes = tb_io.tb_notices(tb)
    assert any(n.startswith("NOT BALANCED") for n in notes) and any(n.startswith("NEGATIVE") for n in notes)
    moved, n = tb_io.move_negatives(tb)
    assert n == 1 and moved.at[0, "credit_2025"] == 100 and moved.at[0, "debit_2025"] == 0


def test_retired_codes_and_grants_rollup():
    from engine import codes
    assert codes.to_new("BS 05.09") == "BS 05.06.01" and codes.to_new("BS 5.09") == "BS 05.06.01"
    assert codes.to_new("PL 2.2.2") == "PL 02.02.01.01" and codes.to_new("PL 8.11") == "PL 08.13.01"
    assert codes.to_new("PL 08.14") == "PL 08.13.01"


# ------------------------------------------------- 2026-09-28 fixed assets / stock / assessment --
def _register():
    from engine import fixed_assets as fa
    return fa.normalise(pd.DataFrame([
        {"description": "Laptop", "category": "BS 01.07", "acq_date": "15/03/2022", "cost": 1_000_000, "doc_no": "EBM1"},
        {"description": "Car", "category": "BS 01.03", "acq_date": "2021-06-01", "cost": 20_000_000, "doc_no": "D1",
         "disp_date": "30/06/2024", "disp_proceeds": 9_000_000},
        {"description": "Building", "category": "BS 01.02", "acq_date": "2020-01-01", "cost": 100_000_000, "doc_no": "B1"}]))


def test_book_straight_line_and_disposal():
    from engine import fixed_assets as fa
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    b = fa.book_schedule(_register(), ch, {}, 2025)
    car = b[b["description"] == "Car"].set_index("year")
    assert list(car["charge"]) == [5e6, 5e6, 5e6, 0.0]                  # full year when bought, none when sold
    assert car.at[2024, "gain_loss"] == 9e6 - 5e6 and car.at[2024, "cost_close"] == 0
    lap = b[b["description"] == "Laptop"].set_index("year")
    assert list(lap["charge"]) == [5e5, 5e5, 0.0, 0.0] and lap.at[2025, "nbv"] == 0
    b2 = fa.book_schedule(_register(), ch, {"disposal_year_full": True}, 2025)
    assert b2[(b2["description"] == "Car") & (b2["year"] == 2024)]["charge"].iloc[0] == 5e6
    life = fa.book_schedule(_register().iloc[[2]], ch, {}, 2025, full_life=True)
    assert life["year"].max() == 2039 and abs(life["dep_close"].iloc[-1] - 1e8) < 1


def test_tax_pools_small_pool_and_disposal():
    from engine import fixed_assets as fa
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    t = fa.tax_schedule(_register(), ch, {}, 2025)
    lap = t[t["description"] == "Laptop"].set_index("year")
    assert lap.at[2022, "allowance"] == 5e5 and lap.at[2023, "rate"] == 1.0 and lap.at[2023, "tax_close"] == 0   # < 500,001
    car = t[t["description"] == "Car"].set_index("year")
    assert car.at[2024, "disposition"] == car.at[2023, "tax_close"] and car.at[2024, "allowance"] == 0   # own tax value out
    bld = t[t["description"] == "Building"].set_index("year")
    assert abs(bld.at[2021, "allowance"] - 95e6 * 0.05) < 1                                                    # reducing balance
    rt = fa.rra_table(t, 2024)
    assert list(rt.columns) == fa.RRA_COLUMNS and len(rt) == 2


def test_register_reconciliation_and_template_roundtrip():
    from engine import fixed_assets as fa
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    reg = _register()
    tb = tb_io.normalise(pd.DataFrame([
        {"code": "BS 01.07", "account": "IT", "debit_2025": 1e6, "credit_2025": 0},
        {"code": "BS 01.02", "account": "Bld", "debit_2025": 1e8, "credit_2025": 0},
        {"code": "BS 01.10", "account": "Acc", "debit_2025": 0, "credit_2025": 1e6 + 3e7},
        {"code": "PL 05.01", "account": "Dep", "debit_2025": 5e6, "credit_2025": 0}]))
    rec = fa.reconcile(fa.book_schedule(reg, ch, {}, 2025), tb, ch, [2025])
    assert fa.year_ok(rec, 2025)
    back = fa.parse_workbook(fa.template_workbook(ch, {}, reg))
    assert len(back) == 3 and back.iloc[1]["disp_proceeds"] == 9e6


def test_stock_form_movement_and_note12():
    from engine import stock
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    form = {"items": [{"item": "A1", "name": "Rice", "category": "BS 03.01.01.04", "unit": "bag"}],
            "opening": [{"item": "A1", "qty": 10, "unit_cost": 100, "value": 1000}],
            "in": [{"item": "A1", "qty": 50, "unit_cost": 100, "value": 5000, "date": "2025-02-01"}],
            "sales": [{"item": "A1", "qty": 40, "unit_cost": 100, "value": 4000}],
            "losses": [{"item": "A1", "qty": 2, "unit_cost": 100, "value": 200, "date": "2025-03-01", "event": "Theft",
                        "doc_no": "PR-1", "comment": ""}],
            "closing": [{"item": "A1", "qty": 17, "unit_cost": 100, "value": 1700}]}
    mv = stock.movement(form, ch)
    assert mv.at[0, "Expected closing qty"] == 18 and mv.at[0, "Qty difference"] == -1
    tb = tb_io.normalise(pd.DataFrame([{"code": "BS 03.01.01.04", "account": "Goods", "debit_2025": 1700, "credit_2025": 0,
                                        "debit_2024": 1000, "credit_2024": 0},
                                       {"code": "PL 02.02.01", "account": "Purch", "debit_2025": 5000, "credit_2025": 0,
                                        "debit_2024": 0, "credit_2024": 0}]))
    assert stock.reconcile(form, tb, ch, 2025)["OK"].all()
    im, _ = stock.note12_from_project({"stock": {"2025": form}}, tb, ch, 2025)
    assert im["opening"] + im["purchases"] + im["cost_of_sales"] + im["gain_loss"] == 1700
    back = stock.parse_workbook(stock.template_workbook(ch, 2025, form))
    assert back["losses"][0]["doc_no"] == "PR-1"


def test_assessment_report(demo):
    from engine import assessment, multiyear
    ch = json.loads(CHART.read_text(encoding="utf-8"))
    proj, tb = multiyear.from_pair_input(demo)
    m = build(demo)
    valid = {l["code"] for p in ("pnl", "bs") for s in ch[p] for g in s.get("groups", []) for l in g["lines"]}
    a = assessment.build(proj, tb, ch, {2025: m}, valid, [], False, "T")
    assert {x["Area"] for x in a["areas"]} >= {"Trial balance", "Taxable income", "Fixed assets", "Stock"}
    assert any("Fixed-asset register" in x for x in a["missing"])
    assert assessment.to_docx(a, True)[:2] == b"PK"
