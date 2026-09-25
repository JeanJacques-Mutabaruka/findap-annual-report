"""Builds the set of files produced by the app.

Files (all prefixed with the same stem  <Company>_Annual_Report_FS__V<yyyy-mm-dd> <hhmm>):
  _input.json       project file — reload it on 1 · Trial Balance to resume work
  _model.json       computed statements, notes and controls (input of the Word renderer)
  _controls.md      controls report (blocking / warnings / info + key figures)
  _statements.xlsx  statements, notes, tax computation, TB with codes and notes, controls
  .docx / .pdf      the annual report (pdf only when LibreOffice is available)
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
BOLD_TYPES = {"heading", "title", "group", "line_bold", "total", "sectiontotal", "subtotal", "grandtotal"}


def company_slug(company: str | None) -> str:
    c = re.sub(r"\b(LTD|LIMITED|PLC|SARL|SA)\b\.?", "", (company or "Company"), flags=re.I).strip()
    return re.sub(r"[^A-Za-z0-9]+", "_", c).strip("_") or "Company"


def stem(company: str | None, when: datetime | None = None, year: int | None = None, kind: str = "Annual_Report_FS") -> str:
    """<Company>_<kind>[_FY2025]__V<yyyy-mm-dd> <hhmm>"""
    when = when or datetime.now()
    fy = f"_FY{year}" if year else ""
    return f"{company_slug(company)}_{kind}{fy}__V{when:%Y-%m-%d %H%M}"


def rows_df(rows: list[dict], cols=("cy", "py"), labels=None) -> pd.DataFrame:
    out = []
    for r in rows:
        rec = {"Line": r["label"], "Note": r.get("note") or ""}
        if "values" in r:
            for i, lab in enumerate(labels or []):
                v = r["values"][i] if r.get("values") else None
                rec[lab] = v
        else:
            for c, lab in zip(cols, labels or cols):
                rec[lab] = r.get(c)
        rec["_type"] = r["type"]
        out.append(rec)
    return pd.DataFrame(out)


def statements_workbook(model: dict, tb: pd.DataFrame | None = None) -> bytes:
    y, p = str(model["meta"]["year_cy"]), str(model["meta"]["year_py"])
    sheets = {
        "P&L": rows_df(model["pnl"], labels=[y, p]),
        "Balance sheet": rows_df(model["bs"], labels=[y, p]),
        "Cash flow": rows_df(model["cashflow"], cols=("cy",), labels=[y]),
        "Income tax": rows_df(model["income_tax"]["rows"], labels=[y, p]),
    }
    eq_rows = []
    for blk in model["equity"]["blocks"]:
        eq_rows.append({"label": blk["title"], "type": "title", "values": [None] * len(model["equity"]["columns"])})
        eq_rows += blk["rows"]
    sheets["Equity"] = rows_df(eq_rows, labels=model["equity"]["columns"])
    if model.get("ppe"):
        sheets["PPE"] = rows_df(model["ppe"]["rows"], labels=model["ppe"]["columns"])
    note_rows = []
    for n in model["notes"]:
        note_rows.append({"Note": n["id"], "Line": n["title"], y: None, p: None, "_type": "title"})
        for b in n["blocks"]:
            if b["kind"] == "text":
                note_rows.append({"Note": "", "Line": b["text"], y: None, p: None, "_type": "text"})
            elif b["kind"] == "table2":
                for r in b["rows"]:
                    note_rows.append({"Note": "", "Line": r["label"], y: r.get("cy"), p: r.get("py"), "_type": r["type"]})
    sheets["Notes"] = pd.DataFrame(note_rows)
    sheets["Controls"] = pd.DataFrame(model["controls"])[["level", "id", "message"]]
    if tb is not None:
        sheets["TB"] = tb.copy()
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        for name, df in sheets.items():
            types = df.pop("_type") if "_type" in df.columns else None
            df.to_excel(xl, sheet_name=name[:31], index=False)
            ws = xl.book[name[:31]]
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.fill, c.font = HEADER_FILL, HEADER_FONT
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for j, col in enumerate(df.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(300).tolist()]) + 2
                ws.column_dimensions[get_column_letter(j)].width = min(max(width, 10), 70)
                if pd.api.types.is_numeric_dtype(df[col]) or col in (y, p) or col in model["equity"]["columns"]:
                    for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2):
                        for c in cell:
                            c.number_format = '#,##0;(#,##0);"-"'
            if types is not None:
                for i, t in enumerate(types.tolist(), start=2):
                    if t in BOLD_TYPES:
                        for c in ws[i]:
                            c.font = Font(bold=True)
        info = pd.DataFrame({"Field": ["Generated at", "Company", "Period end", "Tool"],
                             "Value": [datetime.now().strftime("%d-%b-%Y %H:%M"), model["meta"].get("company_name"),
                                       model["meta"].get("period_end"), "Annual Report Generator V1-0g"]})
        info.to_excel(xl, sheet_name="_INFO", index=False)
    return buf.getvalue()


def bundle(files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


def to_json_bytes(obj) -> bytes:
    return json.dumps(obj, indent=1, ensure_ascii=False, default=str).encode("utf-8")


# ------------------------------------------------------------------ helpers --
def _format_sheet(ws, df: pd.DataFrame, money_cols=()) -> None:
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for j, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(300).tolist()]) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 10), 70)
        if col in money_cols:
            for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2):
                for c in cell:
                    c.number_format = '#,##0;(#,##0);"-"'


def workbook(sheets: dict[str, pd.DataFrame], money: dict[str, list] | None = None, bold_types: dict | None = None) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        for name, df in sheets.items():
            df.to_excel(xl, sheet_name=name[:31], index=False)
            ws = xl.book[name[:31]]
            _format_sheet(ws, df, (money or {}).get(name, []))
            for i, t in enumerate((bold_types or {}).get(name, []), start=2):
                if t in BOLD_TYPES:
                    for c in ws[i]:
                        c.font = Font(bold=True)
    return buf.getvalue()


# ---------------------------------------------------------------- code maps --
def chart_workbook(catalogue: pd.DataFrame) -> bytes:
    """The chart of CIT codes: code -> statement, section, group, line, note."""
    df = catalogue.rename(columns={"code": "CIT code", "statement": "Statement", "section": "Section",
                                   "group": "Group", "line": "Statement line", "note": "Note"})
    return workbook({"CIT codes": df})


def account_map_df(tb: pd.DataFrame, catalogue: pd.DataFrame) -> pd.DataFrame:
    """Account <-> CIT code map of this TB, with the statement line and note each account lands on."""
    cat = catalogue.set_index("code")
    rows = []
    for i, r in tb.iterrows():
        hit = cat.loc[r["code"]] if r["code"] in cat.index else None
        rows.append({"Line": i + 1, "Account": r["account"], "CIT code": r["code"] or "(missing)",
                     "Statement": hit["statement"] if hit is not None else "⚠️ not in chart",
                     "Statement line": hit["line"] if hit is not None else "",
                     "Group": hit["group"] if hit is not None else "",
                     "Note": (r["note"] or (hit["note"] if hit is not None else "")) or ""})
    return pd.DataFrame(rows)


def account_map_workbook(tb: pd.DataFrame, catalogue: pd.DataFrame) -> bytes:
    return workbook({"Account map": account_map_df(tb, catalogue),
                     "CIT codes": catalogue.rename(columns={"code": "CIT code"})})


# ---------------------------------------------------------------- TB template --
TPL_FIXED = ["Statement", "Section", "Group", "Statement line", "CIT code", "Account name"]


def _lists_sheet(wb, catalogue: pd.DataFrame):
    """Hidden sheet 'Lists': every drop-down list as a block of column D, and a key -> range map in A:B.

    Keys: STATEMENTS | ST|<statement> (sections) | SE|<section> (groups) | GR|<group> (lines)
          CST|<statement> / CSE|<section> / CGR|<group> / CLN|<group>|<line> (CIT codes) | ALL (codes)
          ALLSEC / ALLGRP / ALLLIN (full lists when nothing is selected upstream)
    Drop-downs use =INDIRECT(VLOOKUP(key, Lists!$A:$B, 2, FALSE)) — works in every Excel version."""
    ws = wb.create_sheet("Lists")
    blocks: dict[str, list] = {}

    def add(key, values):
        vals = list(dict.fromkeys(v for v in values if v not in (None, "")))  # unique, ordered
        if vals:
            blocks[key] = vals

    cat = catalogue.copy()
    add("STATEMENTS", cat["statement"])
    add("ALLSEC", cat["section"])
    add("ALLGRP", cat["group"])
    add("ALLLIN", cat["line"])
    add("ALL", cat["code"])
    for stt, d in cat.groupby("statement", sort=False):
        add(f"ST|{stt}", d["section"]); add(f"CST|{stt}", d["code"])
    for sec, d in cat.groupby("section", sort=False):
        add(f"SE|{sec}", d["group"]); add(f"CSE|{sec}", d["code"])
    for grp, d in cat.groupby("group", sort=False):
        add(f"GR|{grp}", d["line"]); add(f"CGR|{grp}", d["code"])
    for (grp, line), d in cat.groupby(["group", "line"], sort=False):
        add(f"CLN|{grp}|{line}", d["code"])
    row = 1
    ws.cell(1, 1, "key"); ws.cell(1, 2, "range"); ws.cell(1, 4, "lists")
    r_map = 2
    row = 2
    for key, vals in blocks.items():
        first = row
        for v in vals:
            ws.cell(row, 4, v)
            row += 1
        ws.cell(r_map, 1, key)
        ws.cell(r_map, 2, f"Lists!$D${first}:$D${row - 1}")
        r_map += 1
    # code -> description table used by the check column (F:I)
    ws.cell(1, 6, "code"); ws.cell(1, 7, "statement line"); ws.cell(1, 8, "group"); ws.cell(1, 9, "section")
    for i, r in enumerate(cat.itertuples(), start=2):
        ws.cell(i, 6, r.code); ws.cell(i, 7, r.line); ws.cell(i, 8, r.group); ws.cell(i, 9, r.section)
    ws.sheet_state = "hidden"
    return blocks["STATEMENTS"], len(cat) + 1


def tb_template(years: list[int], catalogue: pd.DataFrame, company: str = "", n_rows: int = 300) -> bytes:
    """Blank trial-balance template.

    Columns: Statement · Section · Group · Statement line · CIT code · Account name · Debit/Credit per year
    (most recent first) · Comments · Check. Cascading drop-downs: choosing a Statement limits the Sections,
    a Section limits the Groups, a Group limits the Statement lines, and the CIT-code list always follows the
    most precise choice made (all codes when nothing is chosen). The Check column shows the statement line of
    the CIT code entered and flags a code that contradicts the line chosen. Re-uploaded, it maps automatically
    (a blank CIT code is derived from Group + Statement line)."""
    from openpyxl import Workbook
    ys = sorted(years, reverse=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Trial Balance"
    _, n_codes = _lists_sheet(wb, catalogue)
    ws["A1"] = f"{company or 'COMPANY NAME'} — TRIAL BALANCE (one line per account)"
    ws["A1"].font = Font(bold=True, size=12, color="1F4E79")
    ws["A2"] = ("Choose Statement › Section › Group › Statement line to narrow the CIT-code list, or pick the CIT code "
                "directly. A code may repeat; each account description must be unique (duplicates turn red).")
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    amount_hdr = [f"{s} {y}" for y in ys for s in ("Debit", "Credit")]
    header = TPL_FIXED + amount_hdr + ["Comments", "Check (statement line of the CIT code)", "_codes"]
    ws.append(header)
    for c in ws[3]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    first, last = 4, 3 + n_rows
    col = {h: get_column_letter(i + 1) for i, h in enumerate(header)}
    A, B, C, D, E = (col[h] for h in TPL_FIXED[:5])
    chk, hlp = col["Check (statement line of the CIT code)"], col["_codes"]
    examples = [("BS 3.1.3.2", "BANK OF KIGALI - RWF"), ("BS 3.1.3.1", "PETTY CASH"), ("BS 3.1.2.1", "ACCOUNTS RECEIVABLE"),
                ("BS 1.08", "COMPUTERS AND OFFICE EQUIPMENT"), ("BS 1.09", "ACCUMULATED DEPRECIATION (credit)"),
                ("BS 8.1.1", "ACCOUNTS PAYABLE"), ("BS 5.01", "SHARE CAPITAL"), ("BS 5.07", "RETAINED EARNINGS (opening)"),
                ("PL 1.1", "SALES"), ("PL 2.2.1", "COST OF SALES - LOCAL PURCHASES"), ("PL 6.01", "STAFF SALARIES"),
                ("PL 5.01", "DEPRECIATION"), ("PL 7.02", "BANK CHARGES")]
    cat = catalogue.set_index("code")
    for i, (code, acc) in enumerate(examples):
        r = first + i
        if code in cat.index:
            h = cat.loc[code]
            ws[f"{A}{r}"], ws[f"{B}{r}"], ws[f"{C}{r}"], ws[f"{D}{r}"] = h["statement"], h["section"], h["group"], h["line"]
        ws[f"{E}{r}"] = code
        ws[f"{col['Account name']}{r}"] = acc
        ws[f"{col['Comments']}{r}"] = "example line — replace or delete"
    lk = "Lists!$A:$B"
    for r in range(first, last + 1):
        # helper: range of the CIT codes allowed by the most precise choice made on the row
        ws[f"{hlp}{r}"] = (f'=IFERROR(VLOOKUP(IF(${D}{r}<>"","CLN|"&${C}{r}&"|"&${D}{r},IF(${C}{r}<>"","CGR|"&${C}{r},'
                           f'IF(${B}{r}<>"","CSE|"&${B}{r},IF(${A}{r}<>"","CST|"&${A}{r},"ALL")))),{lk},2,FALSE),'
                           f'VLOOKUP("ALL",{lk},2,FALSE))')
        ws[f"{chk}{r}"] = (f'=IF(${E}{r}="","",IFERROR(VLOOKUP(${E}{r},Lists!$F$2:$G${n_codes},2,FALSE)&'
                           f'IF(AND(${D}{r}<>"",VLOOKUP(${E}{r},Lists!$F$2:$G${n_codes},2,FALSE)<>${D}{r}),'
                           f'"  ⚠ differs from the line chosen",""),"⚠ unknown CIT code"))')
    def dv(formula, rng, title):
        v = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                           errorTitle=title, error=f"Choose a value from the list ({title}).")
        ws.add_data_validation(v)
        v.add(rng)
    r0 = first
    dv(f'INDIRECT(VLOOKUP("STATEMENTS",{lk},2,FALSE))', f"{A}{first}:{A}{last}", "Statement")
    dv(f'INDIRECT(IFERROR(VLOOKUP("ST|"&${A}{r0},{lk},2,FALSE),VLOOKUP("ALLSEC",{lk},2,FALSE)))', f"{B}{first}:{B}{last}", "Section")
    dv(f'INDIRECT(IFERROR(VLOOKUP("SE|"&${B}{r0},{lk},2,FALSE),VLOOKUP("ALLGRP",{lk},2,FALSE)))', f"{C}{first}:{C}{last}", "Group")
    dv(f'INDIRECT(IFERROR(VLOOKUP("GR|"&${C}{r0},{lk},2,FALSE),VLOOKUP("ALLLIN",{lk},2,FALSE)))', f"{D}{first}:{D}{last}", "Statement line")
    dv(f"INDIRECT(${hlp}{r0})", f"{E}{first}:{E}{last}", "CIT code")
    # a CIT code may repeat, an account description may not: duplicated descriptions turn red
    from openpyxl.formatting.rule import FormulaRule
    acc_c = col["Account name"]
    ws.conditional_formatting.add(
        f"{acc_c}{first}:{acc_c}{last}",
        FormulaRule(formula=[f'AND({acc_c}{first}<>"",COUNTIF(${acc_c}${first}:${acc_c}${last},{acc_c}{first})>1)'],
                    fill=PatternFill("solid", start_color="F8CBAD"), font=Font(color="9C0006", bold=True)))
    widths = {A: 13, B: 22, C: 26, D: 40, E: 12, col["Account name"]: 42, col["Comments"]: 30, chk: 42}
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    ws.column_dimensions[hlp].hidden = True
    for h in amount_hdr:
        L = col[h]
        ws.column_dimensions[L].width = 15
        for row in ws.iter_rows(min_row=first, max_row=last + 1, min_col=ws[L + "1"].column, max_col=ws[L + "1"].column):
            for c in row:
                c.number_format = '#,##0;(#,##0);"-"'
    for c in ws[f"{chk}{first}:{chk}{last}"]:
        c[0].font = Font(size=9, color="7F7F7F")
    ws.freeze_panes = f"{col['Account name']}4"
    tot = last + 1
    ws[f"{col['Account name']}{tot}"] = "TOTAL"
    ws[f"{col['Account name']}{tot}"].font = Font(bold=True)
    for h in amount_hdr:
        L = col[h]
        ws[f"{L}{tot}"] = f"=SUM({L}{first}:{L}{last})"
        ws[f"{L}{tot}"].font = Font(bold=True)
    cs = wb.create_sheet("CIT codes")
    cs.append(["CIT code", "Statement", "Section", "Group", "Statement line", "Note"])
    for r in catalogue.itertuples():
        cs.append([r.code, r.statement, r.section, r.group, r.line, r.note])
    for c in cs[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for L, w in zip("ABCDEF", (12, 14, 26, 30, 60, 6)):
        cs.column_dimensions[L].width = w
    cs.auto_filter.ref = f"A1:F{len(catalogue) + 1}"
    cs.freeze_panes = "A2"
    ins = wb.create_sheet("Instructions", 0)
    lines = [
        "HOW TO FILL THE TRIAL BALANCE TEMPLATE",
        "",
        "1. Sheet 'Trial Balance': one line per ledger account, all years on the same line.",
        f"2. Years in this template: {', '.join(str(y) for y in ys)} (most recent first).",
        "   CY = Current Year (the year being reported); PY = Previous Year (the comparative year before it).",
        "3. Finding the CIT code — two ways:",
        "   a) cascade: choose Statement, then Section, then Group, then Statement line — each list only shows the",
        "      choices that belong to the previous one, and the CIT-code list shows only the matching code(s);",
        "   b) direct: pick the CIT code (all codes are offered when nothing is chosen on the left).",
        "   The grey 'Check' column shows the statement line of the code entered and warns if it contradicts your choice.",
        "   Searching by keyword: sheet 'CIT codes' has a filter on every column (Data > Filter).",
        "4. One CIT code can be used on several lines (e.g. two bank accounts), but EACH LINE NEEDS ITS OWN",
        "   DESCRIPTION: a description used twice turns red — merge the lines or make the descriptions distinct",
        "   (e.g. 'ACCUMULATED DEPRECIATION - VEHICLES' / '… - COMPUTERS').",
        "5. Amounts: positive numbers. Debit balances in 'Debit', credit balances in 'Credit'. Never negative debits.",
        "6. Accumulated depreciation (code BS 1.09) is a CREDIT balance.",
        "7. Retained earnings (BS 5.07) = balance at the START of the year, before the profit of the year.",
        "8. Do not include the income tax charge of the year nor the profit of the year — the generator computes them.",
        "9. No code at all? Leave it empty: the app suggests one from the account name and asks you to confirm.",
        "10. Delete the example lines. The TOTAL line at the bottom must show Debit = Credit for every year.",
        "11. Upload the file on page 1 · Trial Balance — columns and years are recognised automatically.",
        "Note: the cascading lists use standard Excel formulas (INDIRECT); in Google Sheets and some older",
        "LibreOffice versions they may show the full lists instead.",
    ]
    for i, t in enumerate(lines, start=1):
        ins.cell(i, 1, t).font = Font(bold=(i == 1), size=12 if i == 1 else 10, color="1F4E79" if i == 1 else "000000")
    ins.column_dimensions["A"].width = 115
    wb.active = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------ multi-year Excel --
def multi_year_workbook(sheets_rows: dict[str, tuple[list[dict], list[int]]], balance: pd.DataFrame | None = None) -> bytes:
    """sheets_rows: {sheet: (rows with 'values', years)} -> formatted workbook with one column per year."""
    sheets, money, types = {}, {}, {}
    for name, (rows, years) in sheets_rows.items():
        recs = []
        for r in rows:
            rec = {"Line": r["label"], "Note": r.get("note") or ""}
            for y, v in zip(years, r.get("values") or [None] * len(years)):
                rec[str(y)] = v
            recs.append(rec)
        sheets[name] = pd.DataFrame(recs)
        money[name] = [str(y) for y in years]
        types[name] = [r["type"] for r in rows]
    if balance is not None:
        sheets["TB balance check"] = balance
        money["TB balance check"] = ["Total debit", "Total credit", "Difference (debit − credit)"]
    return workbook(sheets, money, types)
