"""Fixed-asset register: book depreciation (straight-line per asset), tax depreciation (RRA reducing balance per
category pool), reconciliation with the trial balance, note 11 figures, RRA "Depreciation Table", Excel files.

Rules (decisions of 28-09-2026):
- Book: straight-line on EACH asset at the category rate, full year in the acquisition year, NO depreciation in the
  disposal year (option: full year), residual value 0, until the cost is fully depreciated.
- Tax (RRA annexure "Depreciation Table"): reducing balance per category (pool), same rates, full year in the
  acquisition year; a disposal removes the asset's OWN tax value (before depreciation of that year); when the pool is
  worth less than small_pool_limit + 1 (Rwf 500,001) after the purchases and disposals of the year and before
  depreciation, the whole pool is written off (100%).
- Rates and the pool limit come from data/chart_of_accounts.json -> fixed_assets (one place), editable per project.
"""
from __future__ import annotations

import io
from datetime import date, datetime

import pandas as pd

TOL = 1.0
MONEY = '#,##0;[Red](#,##0);"-"'
COLS = ["asset_id", "description", "category", "doc_no", "acq_date", "cost", "disp_date", "disp_proceeds", "comment"]
HEADERS = {  # template column -> field. The first three are the RRA "Depreciation Table" columns.
    "Description": "description",
    "Acquisition document Number DMC/EBM Invoice": "doc_no",
    "Acquisition Date(dd/mm/yyyy)": "acq_date",
    "Asset ID (internal code)": "asset_id",
    "Category (CIT code)": "category",
    "Cost (Rwf)": "cost",
    "Disposal date (dd/mm/yyyy)": "disp_date",
    "Disposal proceeds (Rwf)": "disp_proceeds",
    "Comments": "comment",
}
RRA_COLUMNS = ["Description", "Acquisition document Number DMC/EBM Invoice", "Acquisition Date(dd/mm/yyyy)",
               "Book Value Beginning of Period", "Acquisition During the period", "Disposition During the period", "Rate",
               "Depreciation allowance for the Period", "Book Value End of the Period"]


# ------------------------------------------------------------------ settings --
def categories(chart: dict, settings: dict | None = None) -> list[dict]:
    """Categories with the rate in force (project override > chart default)."""
    over = ((settings or {}).get("rates") or {})
    out = []
    for c in chart["fixed_assets"]["categories"]:
        r = over.get(c["code"], c["rate"])
        out.append({**c, "rate": None if r is None or r == "" else float(r)})
    return out


def settings_of(project: dict) -> dict:
    s = dict(project.get("asset_settings") or {})
    s.setdefault("rates", {})
    s.setdefault("disposal_year_full", False)
    s.setdefault("note11", "auto")          # auto (only when reconciled) | always | never
    return s


# ------------------------------------------------------------------ parsing --
def _date(v) -> date | None:
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "NaT", "nan", "None"):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        try:
            return (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(v))).date()
        except (ValueError, OverflowError):
            return None
    txt = str(v).strip()
    try:
        if len(txt) >= 10 and txt[4] == "-" and txt[7] == "-":
            return date.fromisoformat(txt[:10])
        return pd.to_datetime(txt, dayfirst=True).date()
    except (ValueError, TypeError):
        return None


def _num(v) -> float | None:
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "":
        return None
    try:
        return float(str(v).replace(",", "").replace(" ", ""))
    except ValueError:
        return None


def normalise(df: pd.DataFrame) -> pd.DataFrame:
    """Register with the standard columns and types."""
    out = pd.DataFrame(columns=COLS) if df is None else df.copy()
    for c in COLS:
        if c not in out.columns:
            out[c] = None
    out = out[COLS].copy()
    out["acq_date"] = out["acq_date"].apply(_date)
    out["disp_date"] = out["disp_date"].apply(_date)
    for c in ("cost", "disp_proceeds"):
        out[c] = out[c].apply(_num)
    for c in ("asset_id", "description", "category", "doc_no", "comment"):
        out[c] = out[c].apply(lambda v: "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip())
    out["category"] = out["category"].str.upper().str.replace(r"\s+", " ", regex=True)
    keep = (out["description"] != "") | out["cost"].notna()
    return out[keep].reset_index(drop=True)


def parse_workbook(data: bytes) -> pd.DataFrame:
    """Read the sheet 'Asset register' of the template (header row found by its titles)."""
    xl = pd.ExcelFile(io.BytesIO(data))
    sheet = next((s for s in xl.sheet_names if s.lower().startswith("asset register")), xl.sheet_names[0])
    raw = xl.parse(sheet, header=None)
    hdr = next((i for i in range(min(len(raw), 15))
                if {"description", "cost (rwf)"} <= {str(x).strip().lower() for x in raw.iloc[i].tolist()}), None)
    if hdr is None:
        raise ValueError("header row not found — use the template of the app (columns 'Description', 'Cost (Rwf)'…)")
    df = xl.parse(sheet, header=hdr)
    ren = {}
    for col in df.columns:
        k = str(col).strip()
        for h, f in HEADERS.items():
            if k.lower() == h.lower():
                ren[col] = f
    df = df.rename(columns=ren)
    df = df[[c for c in df.columns if c in COLS]]
    reg = normalise(df)
    return reg[~reg["description"].str.upper().str.startswith("EXAMPLE")].reset_index(drop=True)


def issues(reg: pd.DataFrame, chart: dict, settings: dict | None = None) -> pd.DataFrame:
    cats = {c["code"]: c for c in categories(chart, settings)}
    out = []
    today = date.today()
    for i, r in reg.iterrows():
        n = i + 1
        lab = r["description"] or "(no description)"
        if not r["description"]:
            out.append((n, lab, "BLOCKING", "Description missing"))
        if r["category"] not in cats:
            out.append((n, lab, "BLOCKING", f"Category '{r['category']}' is not a fixed-asset CIT code"))
        elif cats[r["category"]]["rate"] is None:
            out.append((n, lab, "WARNING", f"No depreciation rate for {r['category']} {cats[r['category']]['name']} — "
                                           "enter it in ⚙️ Rates & options"))
        if r["acq_date"] is None:
            out.append((n, lab, "BLOCKING", "Acquisition date missing or not a date"))
        elif r["acq_date"] > today:
            out.append((n, lab, "WARNING", "Acquisition date in the future"))
        if r["cost"] is None or r["cost"] <= 0:
            out.append((n, lab, "BLOCKING", "Cost missing, zero or negative"))
        if r["disp_date"] is not None and r["acq_date"] is not None and r["disp_date"] < r["acq_date"]:
            out.append((n, lab, "BLOCKING", "Disposal date before the acquisition date"))
        if not pd.isna(r["disp_proceeds"]) and float(r["disp_proceeds"] or 0) != 0 and r["disp_date"] is None:
            out.append((n, lab, "WARNING", "Disposal proceeds without a disposal date"))
        if not r["doc_no"]:
            out.append((n, lab, "WARNING", "No acquisition document number (DMC / EBM invoice) — required by the RRA "
                                           "Depreciation Table"))
    dup = reg[reg["asset_id"] != ""]["asset_id"]
    for aid in dup[dup.duplicated()].unique():
        out.append(("", aid, "WARNING", f"Asset ID '{aid}' used more than once"))
    return pd.DataFrame(out, columns=["Line", "Asset", "Level", "Issue"])


def valid(reg: pd.DataFrame, chart: dict, settings: dict | None = None) -> pd.DataFrame:
    """Rows usable for the computation (no blocking issue)."""
    cats = {c["code"] for c in categories(chart, settings)}
    ok = reg["category"].isin(cats) & reg["acq_date"].notna() & reg["cost"].fillna(0).gt(0)
    ok &= ~(reg["disp_date"].notna() & reg["acq_date"].notna() &
            (reg["disp_date"].fillna(date.max) < reg["acq_date"].fillna(date.min)))
    return reg[ok].reset_index(drop=True)


# ------------------------------------------------------------ book (straight-line) --
def book_schedule(reg: pd.DataFrame, chart: dict, settings: dict | None = None, last_year: int | None = None,
                  full_life: bool = False) -> pd.DataFrame:
    """One row per asset and year, from the acquisition year to last_year (and to the end of the depreciation when
    full_life): cost and depreciation movements, NBV, gain/loss on disposal."""
    s = settings_of({"asset_settings": settings or {}})
    rates = {c["code"]: c["rate"] for c in categories(chart, s)}
    rows = []
    for i, a in valid(reg, chart, s).iterrows():
        rate, cost = rates[a["category"]], float(a["cost"])
        ya = a["acq_date"].year
        yd = a["disp_date"].year if a["disp_date"] else None
        import math
        end = max(last_year or ya, ya)
        if full_life and rate:
            end = max(end, ya + math.ceil(round(1 / rate, 9)) - 1)
        if yd is not None:
            end = min(end, yd) if not full_life else max(min(end, yd), min(yd, end))
        acc = 0.0
        for y in range(ya, end + 1):
            if yd and y > yd:
                break
            held_open = y > ya
            cost_open = cost if held_open else 0.0
            add = cost if y == ya else 0.0
            dep_open = acc
            charge = 0.0
            if rate and (yd is None or y < yd or (y == yd and s["disposal_year_full"])):
                charge = min(cost * rate, cost - acc)
            acc += charge
            disp = dep_disp = gain = 0.0
            if yd == y:
                disp, dep_disp = -cost, -acc
                gain = (0.0 if pd.isna(a["disp_proceeds"]) else float(a["disp_proceeds"] or 0)) - (cost - acc)
            cost_close = cost_open + add + disp
            dep_close = dep_open + charge + dep_disp
            rows.append({"asset": i, "asset_id": a["asset_id"], "description": a["description"], "category": a["category"],
                         "year": y, "rate": rate, "cost_open": cost_open, "additions": add, "disposals": disp,
                         "cost_close": cost_close, "dep_open": dep_open, "charge": charge, "dep_disposals": dep_disp,
                         "dep_close": dep_close, "nbv": cost_close - dep_close, "gain_loss": gain})
            if yd == y:
                break
    return pd.DataFrame(rows)


# ------------------------------------------------------------- tax (RRA pools) --
def tax_schedule(reg: pd.DataFrame, chart: dict, settings: dict | None = None, last_year: int | None = None) -> pd.DataFrame:
    """One row per asset and year: tax value at the start, acquisition, disposal (own tax value), value before
    depreciation, depreciation allowance (rate of the pool, or 100% when the pool is below the limit), value at the end."""
    s = settings_of({"asset_settings": settings or {}})
    rates = {c["code"]: c["rate"] for c in categories(chart, s)}
    limit = float(chart["fixed_assets"].get("small_pool_limit") or 0)
    reg = valid(reg, chart, s)
    if reg.empty:
        return pd.DataFrame()
    y0 = min(d.year for d in reg["acq_date"])
    y1 = max(last_year or y0, max(d.year for d in reg["acq_date"]))
    tv = {i: 0.0 for i in reg.index}
    rows = []
    for y in range(y0, y1 + 1):
        cur = []
        for i, a in reg.iterrows():
            ya, yd = a["acq_date"].year, (a["disp_date"].year if a["disp_date"] else None)
            if ya > y or (yd is not None and yd < y):
                continue
            opening = tv[i] if ya < y else 0.0
            add = float(a["cost"]) if ya == y else 0.0
            disp = -(opening + add) if yd == y else 0.0
            cur.append([i, a, opening, add, disp, opening + add + disp])
        pools: dict[str, float] = {}
        for i, a, o, ad, dp, before in cur:
            pools[a["category"]] = pools.get(a["category"], 0.0) + before
        for i, a, o, ad, dp, before in cur:
            rate = rates[a["category"]]
            pool = pools[a["category"]]
            written_off = 0 < pool < limit + 1
            if before <= 0 or rate is None:
                dep, used_rate = 0.0, rate
            elif written_off:
                dep, used_rate = before, 1.0
            else:
                dep, used_rate = before * rate, rate
            tv[i] = before - dep
            rows.append({"asset": i, "asset_id": a["asset_id"], "description": a["description"], "category": a["category"],
                         "doc_no": a["doc_no"], "acq_date": a["acq_date"], "year": y, "tax_open": o, "acquisition": ad,
                         "disposition": -dp, "before_dep": before, "pool_before_dep": pool, "rate": used_rate,
                         "pool_written_off": written_off, "allowance": dep, "tax_close": tv[i]})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- summaries --
def summary(book: pd.DataFrame, chart: dict) -> pd.DataFrame:
    """Per year and category: cost and depreciation movements, NBV, gain/loss."""
    if book.empty:
        return pd.DataFrame()
    names = {c["code"]: c["name"] for c in chart["fixed_assets"]["categories"]}
    num = ["cost_open", "additions", "disposals", "cost_close", "dep_open", "charge", "dep_disposals", "dep_close", "nbv",
           "gain_loss"]
    g = book.groupby(["year", "category"], as_index=False)[num].sum()
    g.insert(2, "name", g["category"].map(names))
    return g.sort_values(["year", "category"]).reset_index(drop=True)


def tax_summary(tax: pd.DataFrame, chart: dict) -> pd.DataFrame:
    if tax.empty:
        return pd.DataFrame()
    names = {c["code"]: c["name"] for c in chart["fixed_assets"]["categories"]}
    g = tax.groupby(["year", "category"], as_index=False).agg(
        tax_open=("tax_open", "sum"), acquisition=("acquisition", "sum"), disposition=("disposition", "sum"),
        before_dep=("before_dep", "sum"), allowance=("allowance", "sum"), tax_close=("tax_close", "sum"),
        written_off=("pool_written_off", "max"), rate=("rate", "max"))
    g.insert(2, "name", g["category"].map(names))
    return g.sort_values(["year", "category"]).reset_index(drop=True)


def rra_table(tax: pd.DataFrame, year: int) -> pd.DataFrame:
    """RRA annexure sheet 'Depreciation Table' for one year: one row per asset, the 9 RRA columns."""
    t = tax[tax["year"] == year] if not tax.empty else tax
    rows = []
    for _, r in t.sort_values(["category", "acq_date"]).iterrows():
        if max(abs(r["tax_open"]), abs(r["acquisition"]), abs(r["disposition"]), abs(r["allowance"]), abs(r["tax_close"])) <= 0.5:
            continue
        rows.append({"Description": f"{r['description']} [{r['category']}]", RRA_COLUMNS[1]: r["doc_no"],
                     RRA_COLUMNS[2]: r["acq_date"].strftime("%d/%m/%Y"), RRA_COLUMNS[3]: round(r["tax_open"], 0),
                     RRA_COLUMNS[4]: round(r["acquisition"], 0), RRA_COLUMNS[5]: round(r["disposition"], 0),
                     "Rate": r["rate"], RRA_COLUMNS[7]: round(r["allowance"], 0), RRA_COLUMNS[8]: round(r["tax_close"], 0)})
    return pd.DataFrame(rows, columns=RRA_COLUMNS)


# ----------------------------------------------------------- TB reconciliation --
def _tb_net(tb: pd.DataFrame, codes: list[str], y: int) -> float:
    if tb is None or tb.empty or f"debit_{y}" not in tb.columns:
        return 0.0
    sel = tb[tb["code"].isin(codes)]
    return float(sel[f"debit_{y}"].sum() - sel[f"credit_{y}"].sum())


def reconcile(book: pd.DataFrame, tb: pd.DataFrame, chart: dict, years: list[int]) -> pd.DataFrame:
    """Register vs TB for each year: cost per category, accumulated depreciation, charge, disposal result, NBV."""
    names = {c["code"]: c["name"] for c in chart["fixed_assets"]["categories"]}
    rows = []
    for y in years:
        b = book[book["year"] == y] if not book.empty else book
        for code, name in names.items():
            reg_v = float(b[b["category"] == code]["cost_close"].sum()) if not b.empty else 0.0
            tb_v = _tb_net(tb, [code], y)
            if abs(reg_v) > TOL or abs(tb_v) > TOL:
                rows.append({"Year": y, "Item": f"Cost — {name}", "Code": code, "Register": reg_v, "Trial balance": tb_v})
        reg_dep = float(b["dep_close"].sum()) if not b.empty else 0.0
        rows.append({"Year": y, "Item": "Accumulated depreciation (all categories)", "Code": "BS 01.10",
                     "Register": reg_dep, "Trial balance": -_tb_net(tb, ["BS 01.10"], y)})
        rows.append({"Year": y, "Item": "Depreciation charge of the year", "Code": "PL 05.01 + PL 02.04.06",
                     "Register": float(b["charge"].sum()) if not b.empty else 0.0,
                     "Trial balance": _tb_net(tb, ["PL 05.01", "PL 02.04.06"], y)})
        rows.append({"Year": y, "Item": "Gain / (loss) on disposals", "Code": "PL 08.11 − PL 05.02",
                     "Register": float(b["gain_loss"].sum()) if not b.empty else 0.0,
                     "Trial balance": -_tb_net(tb, ["PL 08.11", "PL 05.02"], y)})
        cost_codes = list(names)
        rows.append({"Year": y, "Item": "Net book value", "Code": "BS 01.xx",
                     "Register": (float(b["nbv"].sum()) if not b.empty else 0.0),
                     "Trial balance": _tb_net(tb, cost_codes + ["BS 01.10"], y)})
    df = pd.DataFrame(rows)
    if not df.empty:
        df["Difference"] = df["Register"] - df["Trial balance"]
        df["OK"] = df["Difference"].abs() <= TOL
    return df


def year_ok(rec: pd.DataFrame, y: int) -> bool:
    r = rec[rec["Year"] == y]
    return bool(len(r)) and bool(r["OK"].all())


def note11_categories(book: pd.DataFrame, chart: dict, year: int) -> list[dict] | None:
    """Categories for note 11 (build_model fixed_assets.categories) from the register."""
    s = summary(book, chart)
    if s.empty:
        return None
    s = s[s["year"] == year]
    if s.empty:
        return None
    return [{"name": r["name"], "cost_opening": r["cost_open"], "additions": r["additions"], "disposals": r["disposals"],
             "dep_opening": r["dep_open"], "charge": r["charge"], "dep_on_disposals": r["dep_disposals"]}
            for _, r in s.iterrows() if abs(r["cost_open"]) + abs(r["additions"]) + abs(r["cost_close"]) > TOL]


def book_vs_tax(book: pd.DataFrame, tax: pd.DataFrame, years: list[int]) -> pd.DataFrame:
    rows = []
    for y in years:
        bk = float(book[book["year"] == y]["charge"].sum()) if not book.empty else 0.0
        tx = float(tax[tax["year"] == y]["allowance"].sum()) if not tax.empty else 0.0
        rows.append({"Year": y, "Book depreciation (straight-line)": bk, "Tax depreciation (RRA pools)": tx,
                     "Difference (tax − book)": tx - bk})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ Excel files --
def _st():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    return {"h": Font(name="Arial", bold=True, color="FFFFFF", size=10), "hf": PatternFill("solid", fgColor="1F4E79"),
            "n": Font(name="Arial", size=10), "b": Font(name="Arial", bold=True, size=10),
            "inp": Font(name="Arial", size=10, color="0000FF"), "inpf": PatternFill("solid", fgColor="FFF2CC"),
            "grey": PatternFill("solid", fgColor="EDEDED"), "tot": PatternFill("solid", fgColor="DDEBF7"),
            "rraf": PatternFill("solid", fgColor="E2EFDA"), "wrap": Alignment(wrap_text=True, vertical="top"),
            "bd": Border(top=Side(style="thin", color="7F7F7F")), "i": Font(name="Arial", italic=True, size=9, color="595959"),
            "title": Font(name="Arial", bold=True, size=13, color="1F4E79")}


def template_workbook(chart: dict, settings: dict | None = None, reg: pd.DataFrame | None = None) -> bytes:
    """Asset-register template (RRA Depreciation Table columns first + the columns the app needs), with the list of
    categories and rates, a drop-down on the category, an example row and the instructions."""
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    S = _st()
    cats = categories(chart, settings)
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    lines = ["FIXED-ASSET REGISTER — HOW TO FILL IT", "",
             "1. Sheet 'Asset register': one line per asset (one computer = one line; identical items bought on the same "
             "invoice may share a line).",
             "2. Green columns = columns of the RRA CIT annexure 'Depreciation Table' (Description, DMC/EBM invoice number, "
             "acquisition date). The other columns of the RRA table (book value at the beginning, acquisitions, "
             "dispositions, rate, allowance, book value at the end) are COMPUTED by the app for each year.",
             "3. Category: choose the CIT code in the list (sheet 'Categories & rates'). Computers and accessories = BS 01.07.",
             "4. Acquisition date dd/mm/yyyy and COST (amount paid, excluding recoverable VAT). No depreciation to enter: "
             "the app computes it.",
             "5. Sold / scrapped assets: keep the line and fill the disposal date and the proceeds (0 if scrapped).",
             "6. Asset ID: your internal code if any (useful to find an asset; optional).",
             "7. Delete the EXAMPLE line, save, and upload the file on page 8 · Long-term Assets Check.",
             "",
             "Book depreciation: straight-line on each asset at the category rate, full year in the acquisition year, none in "
             "the disposal year (option), residual value 0.",
             "Tax depreciation (RRA): reducing balance per category at the same rates; a category worth less than Rwf "
             f"{float(chart['fixed_assets'].get('small_pool_limit') or 0) + 1:,.0f} after the purchases and disposals of "
             "the year is written off at 100%."]
    for i, t in enumerate(lines, 1):
        ins.cell(i, 1, t).font = S["title"] if i == 1 else S["n"]
        ins.cell(i, 1).alignment = S["wrap"]
    ins.column_dimensions["A"].width = 120
    ws = wb.create_sheet("Asset register")
    hdr = list(HEADERS)
    ws.append(hdr)
    for j, c in enumerate(ws[1], 1):
        c.font, c.alignment = S["h"], S["wrap"]
        c.fill = S["hf"]
        if j <= 3:
            c.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor="548235")
    data = reg if reg is not None and not reg.empty else None
    if data is None:
        ws.append(["EXAMPLE — Laptop HP ProBook (delete this line)", "EBM 123456", "15/03/2024", "IT-001", "BS 01.07",
                   950000, "", "", ""])
    else:
        for _, r in data.iterrows():
            ws.append([r["description"], r["doc_no"], r["acq_date"].strftime("%d/%m/%Y") if r["acq_date"] else "",
                       r["asset_id"], r["category"], r["cost"], r["disp_date"].strftime("%d/%m/%Y") if r["disp_date"] else "",
                       r["disp_proceeds"], r["comment"]])
    n = max(ws.max_row, 500)
    for rr in range(2, n + 1):
        for col in (6, 8):
            ws.cell(rr, col).number_format = "#,##0"
    for col, w in zip("ABCDEFGHI", (42, 22, 16, 16, 16, 15, 16, 15, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"
    cs = wb.create_sheet("Categories & rates")
    cs.append(["CIT code", "Category", "Rate (book and tax)"])
    for c in cs[1]:
        c.font, c.fill = S["h"], S["hf"]
    for c in cats:
        cs.append([c["code"], c["name"], c["rate"] if c["rate"] is not None else "rate to enter in the app"])
        cs.cell(cs.max_row, 3).number_format = "0%"
    for col, w in zip("ABC", (12, 44, 22)):
        cs.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1=f"='Categories & rates'!$A$2:$A${len(cats) + 1}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:E{n}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_workbook(reg: pd.DataFrame, chart: dict, settings: dict | None, years: list[int], tb: pd.DataFrame | None,
                    company: str = "") -> bytes:
    """Register, book depreciation grid (live straight-line formulas), tax grid (live pool formulas), summaries per
    category and year (SUMIFS), reconciliation with the TB and the RRA Depreciation Table of the last year."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter as L
    S = _st()
    s = settings_of({"asset_settings": settings or {}})
    cats = categories(chart, s)
    reg = valid(reg, chart, s)
    limit = float(chart["fixed_assets"].get("small_pool_limit") or 0)
    y0 = min([d.year for d in reg["acq_date"]] + years) if len(reg) else min(years)
    y1 = max(years) if years else y0
    YRS = list(range(y0, y1 + 1))
    wb = Workbook()
    rm = wb.active
    rm.title = "Read me"
    for i, t in enumerate([f"{company} — Fixed-asset register and depreciation ({y0}–{y1})",
                           f"Generated {datetime.now():%d-%m-%Y %H:%M}. Blue on yellow = inputs (register, rates, options); "
                           "every other figure is a formula.",
                           "Book depreciation: straight-line per asset, full year in the acquisition year, "
                           + ("full year" if s["disposal_year_full"] else "none") + " in the disposal year, residual 0.",
                           f"Tax depreciation: reducing balance per category; pool < {limit + 1:,.0f} before depreciation "
                           "→ 100%; a disposal removes the asset's own tax value.",
                           "Sheets: Register · Rates · Book depreciation · Book accumulated · Tax value before dep · "
                           "Tax allowance · Summary by category · Reconciliation TB · RRA Depreciation Table."], 1):
        rm.cell(i, 1, t).font = S["title"] if i == 1 else S["n"]
    rm.column_dimensions["A"].width = 130
    # Rates
    ra = wb.create_sheet("Rates")
    ra.append(["CIT code", "Category", "Rate"])
    for c in ra[1]:
        c.font, c.fill = S["h"], S["hf"]
    for c in cats:
        ra.append([c["code"], c["name"], c["rate"] if c["rate"] is not None else 0])
        ra.cell(ra.max_row, 3).number_format = "0%"
        ra.cell(ra.max_row, 3).font, ra.cell(ra.max_row, 3).fill = S["inp"], S["inpf"]
    ra.append([])
    ra.append(["Small pool limit (tax)", "", limit])
    ra.append(["Depreciation in the disposal year (1 = full year, 0 = none)", "", 1 if s["disposal_year_full"] else 0])
    lim_ref, dfull_ref = f"Rates!$C${len(cats) + 3}", f"Rates!$C${len(cats) + 4}"
    for rr in (len(cats) + 3, len(cats) + 4):
        ra.cell(rr, 3).font, ra.cell(rr, 3).fill = S["inp"], S["inpf"]
    ra.cell(len(cats) + 3, 3).number_format = "#,##0"
    for col, w in zip("ABC", (50, 44, 12)):
        ra.column_dimensions[col].width = w
    # Register (inputs + helper columns)
    rg = wb.create_sheet("Register")
    rg.append(["Asset ID", "Description", "Category", "DMC/EBM invoice", "Acquisition date", "Cost", "Disposal date",
               "Disposal proceeds", "Acq. year", "Disposal year", "Rate"])
    for c in rg[1]:
        c.font, c.fill, c.alignment = S["h"], S["hf"], S["wrap"]
    for i, a in reg.iterrows():
        r = i + 2
        rg.append([a["asset_id"], a["description"], a["category"], a["doc_no"], a["acq_date"], float(a["cost"]),
                   a["disp_date"], 0.0 if pd.isna(a["disp_proceeds"]) else float(a["disp_proceeds"] or 0), f"=YEAR(E{r})", f'=IF(G{r}="","",YEAR(G{r}))',
                   f"=IFERROR(VLOOKUP(C{r},Rates!$A:$C,3,FALSE),0)"])
        for col in range(1, 9):
            rg.cell(r, col).font, rg.cell(r, col).fill = S["inp"], S["inpf"]
        rg.cell(r, 5).number_format = rg.cell(r, 7).number_format = "dd/mm/yyyy"
        rg.cell(r, 6).number_format = rg.cell(r, 8).number_format = MONEY
        rg.cell(r, 11).number_format = "0%"
    nr = len(reg) + 1
    for col, w in zip("ABCDEFGHIJK", (12, 40, 12, 16, 13, 15, 13, 15, 9, 9, 7)):
        rg.column_dimensions[col].width = w
    rg.freeze_panes = "C2"

    def grid(title, formula, note):
        g = wb.create_sheet(title)
        g.append(["Asset ID", "Description", "Category"] + YRS)
        for c in g[1]:
            c.font, c.fill = S["h"], S["hf"]
        g.cell(2, 1, note).font = S["i"]
        for i in range(len(reg)):
            r, rr = i + 3, i + 2      # grid row, register row
            g.cell(r, 1, f'=IF(Register!A{rr}="","",Register!A{rr})'); g.cell(r, 2, f"=Register!B{rr}"); g.cell(r, 3, f"=Register!C{rr}")
            for j, y in enumerate(YRS):
                col = 4 + j
                g.cell(r, col, formula(r, rr, col, j)).number_format = MONEY
        for col, w in zip("ABC", (12, 36, 12)):
            g.column_dimensions[col].width = w
        for j in range(len(YRS)):
            g.column_dimensions[L(4 + j)].width = 13
        g.freeze_panes = "D3"
        return g

    # Book charge: straight-line, capped at the remaining cost
    def f_charge(r, rr, col, j):
        y = f"{L(col)}$1"
        prev = f"SUM($D{r}:{L(col - 1)}{r})" if j > 0 else "0"
        held = (f"AND({y}>=Register!$I{rr},OR(Register!$J{rr}=\"\",{y}<Register!$J{rr},"
                f"AND({y}=Register!$J{rr},{dfull_ref}=1)))")
        return f"=IF({held},MIN(Register!$F{rr}*Register!$K{rr},Register!$F{rr}-{prev}),0)"
    grid("Book depreciation", f_charge, "Charge of the year = MIN(cost × rate, cost − charges of the previous years) while held")

    def f_acc(r, rr, col, j):
        y = f"{L(col)}$1"
        return (f"=IF(OR({y}<Register!$I{rr},AND(Register!$J{rr}<>\"\",{y}>=Register!$J{rr})),0,"
                f"SUM('Book depreciation'!$D{r}:{L(col)}{r}))")
    grid("Book accumulated", f_acc, "Accumulated depreciation at the year end (0 after the disposal)")

    def f_before(r, rr, col, j):
        y = f"{L(col)}$1"
        prev = (f"('Tax value before dep'!{L(col - 1)}{r}-'Tax allowance'!{L(col - 1)}{r})" if j > 0 else "0")
        return (f"=IF(OR({y}<Register!$I{rr},AND(Register!$J{rr}<>\"\",{y}>=Register!$J{rr})),0,"
                f"IF({y}=Register!$I{rr},Register!$F{rr},{prev}))")
    grid("Tax value before dep", f_before, "Tax value after purchases and disposals, before depreciation (0 once disposed)")

    def f_allow(r, rr, col, j):
        pool = f"SUMIF($C$3:$C${len(reg) + 2},$C{r},'Tax value before dep'!{L(col)}$3:{L(col)}${len(reg) + 2})"
        v = f"'Tax value before dep'!{L(col)}{r}"
        return f"=IF({v}<=0,0,IF({pool}<{lim_ref}+1,{v},{v}*Register!$K{rr}))"
    grid("Tax allowance", f_allow, "Allowance = value × rate, or 100% when the category pool is below the limit + 1")

    # Summary by category and year (SUMIF on the grids)
    sm = wb.create_sheet("Summary by category")
    sm.append(["Year", "CIT code", "Category", "Cost at year end", "Charge of the year", "Accumulated depreciation",
               "Net book value", "Tax allowance", "Tax value at year end"])
    for c in sm[1]:
        c.font, c.fill, c.alignment = S["h"], S["hf"], S["wrap"]
    last = len(reg) + 2
    rng = lambda sh, col: f"'{sh}'!{col}3:{col}{last}"
    for j, y in enumerate(YRS):
        col = L(4 + j)
        for c in cats:
            r = sm.max_row + 1
            cost = (f"SUMPRODUCT((Register!$C$2:$C${nr}=B{r})*(Register!$I$2:$I${nr}<=A{r})*"
                    f"((Register!$J$2:$J${nr}=\"\")+(Register!$J$2:$J${nr}>A{r})>0)*Register!$F$2:$F${nr})")
            sm.append([y, c["code"], c["name"], f"={cost}",
                       f"=SUMIF('Book depreciation'!$C$3:$C${last},B{r},{rng('Book depreciation', col)})",
                       f"=SUMIF('Book accumulated'!$C$3:$C${last},B{r},{rng('Book accumulated', col)})",
                       f"=D{r}-F{r}",
                       f"=SUMIF('Tax allowance'!$C$3:$C${last},B{r},{rng('Tax allowance', col)})",
                       f"=SUMIF('Tax value before dep'!$C$3:$C${last},B{r},{rng('Tax value before dep', col)})-H{r}"])
            for cc in sm[r][3:]:
                cc.number_format = MONEY
    for col, w in zip("ABCDEFGHI", (7, 11, 40, 16, 16, 16, 16, 16, 16)):
        sm.column_dimensions[col].width = w
    sm.freeze_panes = "D2"
    sm.auto_filter.ref = f"A1:I{sm.max_row}"
    # Reconciliation (values)
    if tb is not None and years:
        book = book_schedule(reg, chart, s, max(years))
        rec = reconcile(book, tb, chart, years)
        rc = wb.create_sheet("Reconciliation TB")
        rc.append(list(rec.columns))
        for c in rc[1]:
            c.font, c.fill = S["h"], S["hf"]
        red = __import__("openpyxl").styles.Font(name="Arial", size=10, color="C00000", bold=True)
        for rw in rec.itertuples(index=False):
            rc.append(["" if pd.isna(v) else v for v in rw])
            for cc in rc[rc.max_row][3:6]:
                cc.number_format = MONEY
            if not rw.OK:
                for cc in rc[rc.max_row]:
                    cc.font = red
        for col, w in zip("ABCDEFG", (7, 44, 22, 16, 16, 16, 6)):
            rc.column_dimensions[col].width = w
    # RRA table of the last year
    tax = tax_schedule(reg, chart, s, y1)
    rt = wb.create_sheet(f"RRA Depreciation Table {y1}")
    t = rra_table(tax, y1)
    rt.append(RRA_COLUMNS)
    for c in rt[1]:
        c.font, c.fill, c.alignment = S["h"], S["hf"], S["wrap"]
    for rw in t.itertuples(index=False):
        rt.append(list(rw))
        for cc in rt[rt.max_row][3:9]:
            cc.number_format = MONEY
        rt.cell(rt.max_row, 7).number_format = "0%"
    for col, w in zip("ABCDEFGHI", (44, 20, 14, 16, 16, 16, 7, 16, 16)):
        rt.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def register_from_records(records: list[dict] | None) -> pd.DataFrame:
    return normalise(pd.DataFrame(records or [], columns=COLS))


def register_to_records(reg: pd.DataFrame) -> list[dict]:
    out = []
    for _, r in reg.iterrows():
        d = {k: r[k] for k in COLS}
        d["acq_date"] = r["acq_date"].isoformat() if r["acq_date"] else None
        d["disp_date"] = r["disp_date"].isoformat() if r["disp_date"] else None
        out.append(d)
    return out


def note11_from_project(project: dict, tb: pd.DataFrame, chart: dict, year: int) -> tuple[list | None, str]:
    """Categories of note 11 for CY = year from the register, and why (used / not used)."""
    reg = register_from_records(project.get("asset_register"))
    if reg.empty:
        return None, "no fixed-asset register — note 11 allocates the TB depreciation pro rata to cost"
    s = settings_of(project)
    if s["note11"] == "never":
        return None, "register not used for note 11 (option)"
    book = book_schedule(reg, chart, s, year)
    ys = [y for y in (year - 1, year) if f"debit_{y}" in tb.columns]
    rec = reconcile(book, tb, chart, ys)
    ok = all(year_ok(rec, y) for y in ys)
    if s["note11"] == "auto" and not ok:
        return None, ("register NOT used for note 11: it does not reconcile with the TB for " +
                      ", ".join(str(y) for y in ys if not year_ok(rec, y)) + " (8 · Long-term Assets Check → ⚖️ Reconciliation)")
    cats = note11_categories(book, chart, year)
    return cats, "note 11 built from the fixed-asset register" + ("" if ok else " — WITH DIFFERENCES from the TB (option)")
