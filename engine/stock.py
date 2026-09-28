"""Stock form: item level, grouped by the RRA inventory categories (BS 03.01.01.01 – .04).

One form per financial year: opening stock, movements IN (purchases), movements OUT (sales at cost, and losses /
damages with their detail), closing stock counted. The app computes, per item and category, the expected closing
stock (opening + in − sales − losses), compares it with the count and reconciles the totals with the trial balance.
"""
from __future__ import annotations

import io
from datetime import date, datetime

import pandas as pd

TOL = 1.0
MONEY = '#,##0;[Red](#,##0);"-"'
SHEETS = {  # sheet -> [(column title, field)]
    "Items": [("Item code *", "item"), ("Item name *", "name"), ("Category (CIT code) *", "category"), ("Unit", "unit")],
    "Opening stock": [("Item code *", "item"), ("Quantity", "qty"), ("Unit cost (Rwf)", "unit_cost"), ("Value (Rwf)", "value")],
    "IN - Purchases": [("Date", "date"), ("Item code *", "item"), ("Quantity", "qty"), ("Unit cost (Rwf)", "unit_cost"),
                       ("Value (Rwf)", "value"), ("Supplier", "party"), ("Document number", "doc_no")],
    "OUT - Sales": [("Date", "date"), ("Item code *", "item"), ("Quantity", "qty"), ("Unit cost (Rwf)", "unit_cost"),
                    ("Value at cost (Rwf)", "value"), ("Document number", "doc_no")],
    "OUT - Losses & damages": [("Date of the event", "date"), ("Item code *", "item"), ("Quantity", "qty"),
                               ("Unit cost (Rwf)", "unit_cost"), ("Value (Rwf)", "value"), ("Event", "event"),
                               ("Supporting document number", "doc_no"), ("Comments", "comment")],
    "Closing stock": [("Item code *", "item"), ("Quantity counted", "qty"), ("Unit cost (Rwf)", "unit_cost"),
                      ("Value (Rwf)", "value")],
}
KEYS = {"Items": "items", "Opening stock": "opening", "IN - Purchases": "in", "OUT - Sales": "sales",
        "OUT - Losses & damages": "losses", "Closing stock": "closing"}


def cats(chart: dict) -> list[dict]:
    return chart["stock"]["categories"]


# ------------------------------------------------------------------ template --
def template_workbook(chart: dict, year: int, data: dict | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    h, hf = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    lines = [f"STOCK FORM — financial year {year}", "",
             "1. 'Items': one line per stock item — code, name, category (list) and unit. Use the SAME item code in every sheet.",
             "2. 'Opening stock': quantity and unit cost of each item at the START of the year (= closing stock of last year).",
             "3. 'IN - Purchases': every purchase (or a total per item for the year): quantity, unit cost excluding VAT.",
             "4. 'OUT - Sales': quantities sold, valued AT COST (not at the selling price).",
             "5. 'OUT - Losses & damages': every loss (theft, damage, expiry, fire or other unforeseen event) with its DATE, "
             "the SUPPORTING DOCUMENT number (police report, write-off minutes…) and a comment.",
             "6. 'Closing stock': the quantities COUNTED at the year end and their unit cost.",
             "7. The Value columns are formulas (quantity × unit cost) — you may overwrite them with the real value.",
             "8. Upload the file on page 9 · Stock Control for the year.",
             "",
             "Categories = RRA CIT annexure lines 3.1.1.1 to 3.1.1.4 (balance sheet inventories)."]
    for i, t in enumerate(lines, 1):
        ins.cell(i, 1, t).font = Font(name="Arial", bold=(i == 1), size=13 if i == 1 else 10, color="1F4E79" if i == 1 else "000000")
    ins.column_dimensions["A"].width = 120
    cl = wb.create_sheet("Categories")
    cl.append(["CIT code", "Category"])
    for c in cl[1]:
        c.font, c.fill = h, hf
    for c in cats(chart):
        cl.append([c["code"], c["name"]])
    cl.append([])
    cl.append(["Loss events"])
    ev_start = cl.max_row + 1
    for e in chart["stock"]["loss_events"]:
        cl.append([e])
    ev_end = cl.max_row
    cl.column_dimensions["A"].width = 18
    cl.column_dimensions["B"].width = 44
    data = data or {}
    for sheet, cols in SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append([t for t, _ in cols])
        for c in ws[1]:
            c.font, c.fill, c.alignment = h, hf, Alignment(wrap_text=True, vertical="top")
        fields = [f for _, f in cols]
        rows = data.get(KEYS[sheet]) or []
        for r in rows:
            ws.append([r.get(f) for f in fields])
        n = max(ws.max_row + 200, 300)
        if "value" in fields and "qty" in fields and "unit_cost" in fields:
            vc, qc, uc = (fields.index(x) + 1 for x in ("value", "qty", "unit_cost"))
            from openpyxl.utils import get_column_letter as L
            for rr in range(2, n + 1):
                if ws.cell(rr, vc).value in (None, ""):
                    ws.cell(rr, vc, f'=IF(OR({L(qc)}{rr}="",{L(uc)}{rr}=""),"",{L(qc)}{rr}*{L(uc)}{rr})')
                for cc in (qc, uc, vc):
                    ws.cell(rr, cc).number_format = "#,##0.##" if cc == qc else "#,##0"
        if "date" in fields:
            for rr in range(2, n + 1):
                ws.cell(rr, fields.index("date") + 1).number_format = "dd/mm/yyyy"
        if sheet == "Items":
            dv = DataValidation(type="list", formula1=f"=Categories!$A$2:$A${len(cats(chart)) + 1}", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"C2:C{n}")
        if sheet == "OUT - Losses & damages":
            dv = DataValidation(type="list", formula1=f"=Categories!$A${ev_start}:$A${ev_end}", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"F2:F{n}")
        for j in range(1, len(cols) + 1):
            ws.column_dimensions[chr(64 + j)].width = 34 if cols[j - 1][1] in ("name", "comment", "party") else 16
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------- parsing --
def _num(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "":
        return None
    try:
        return float(str(v).replace(",", "").replace(" ", ""))
    except ValueError:
        return None


def _date(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "NaT"):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    try:
        t = str(v).strip()
        d = date.fromisoformat(t[:10]) if len(t) >= 10 and t[4] == "-" else pd.to_datetime(t, dayfirst=True).date()
        return d.isoformat()
    except (ValueError, TypeError):
        return None


def parse_workbook(data: bytes) -> dict:
    """{"items": [...], "opening": [...], "in": [...], "sales": [...], "losses": [...], "closing": [...]}"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True)
    out = {}
    for sheet, cols in SHEETS.items():
        ws = next((wb[s] for s in wb.sheetnames if s.strip().lower() == sheet.lower()), None)
        rows = []
        if ws is not None:
            hdr = [str(c.value or "").strip().lower() for c in ws[1]]
            idx = {f: (hdr.index(t.lower()) if t.lower() in hdr else None) for t, f in cols}
            for r in ws.iter_rows(min_row=2, values_only=True):
                rec = {f: (r[i] if i is not None and i < len(r) else None) for f, i in idx.items()}
                if not str(rec.get("item") or "").strip():
                    continue
                rec["item"] = str(rec["item"]).strip()
                for f in ("qty", "unit_cost", "value"):
                    if f in rec:
                        rec[f] = _num(rec[f])
                if "value" in rec and rec["value"] is None and rec.get("qty") is not None and rec.get("unit_cost") is not None:
                    rec["value"] = rec["qty"] * rec["unit_cost"]
                if "date" in rec:
                    rec["date"] = _date(rec["date"])
                for f in ("name", "category", "unit", "party", "doc_no", "event", "comment"):
                    if f in rec:
                        rec[f] = "" if rec[f] is None else str(rec[f]).strip()
                if "category" in rec:
                    rec["category"] = rec["category"].upper()
                rows.append(rec)
        out[KEYS[sheet]] = rows
    return out


def issues(form: dict, chart: dict) -> pd.DataFrame:
    codes = {c["code"] for c in cats(chart)}
    items = {r["item"]: r for r in form.get("items", [])}
    out = []
    for r in form.get("items", []):
        if r.get("category") not in codes:
            out.append(("Items", r["item"], "BLOCKING", f"Category '{r.get('category')}' is not an inventory CIT code"))
    for key, sheet in (("opening", "Opening stock"), ("in", "IN - Purchases"), ("sales", "OUT - Sales"),
                       ("losses", "OUT - Losses & damages"), ("closing", "Closing stock")):
        for r in form.get(key, []):
            if r["item"] not in items:
                out.append((sheet, r["item"], "BLOCKING", "Item code not in the sheet 'Items'"))
            if r.get("value") is None:
                out.append((sheet, r["item"], "WARNING", "No value (quantity × unit cost missing)"))
            if key == "losses":
                if not r.get("date"):
                    out.append((sheet, r["item"], "WARNING", "Loss without date"))
                if not r.get("doc_no"):
                    out.append((sheet, r["item"], "WARNING", "Loss without supporting document number"))
                if not r.get("event"):
                    out.append((sheet, r["item"], "WARNING", "Loss without event type"))
    return pd.DataFrame(out, columns=["Sheet", "Item", "Level", "Issue"])


# --------------------------------------------------------------- computation --
def movement(form: dict, chart: dict) -> pd.DataFrame:
    """Per item: opening, in, sales, losses, expected closing, counted closing, difference (quantities and values)."""
    items = {r["item"]: r for r in form.get("items", [])}
    names = {c["code"]: c["name"] for c in cats(chart)}
    agg = {}
    for key in ("opening", "in", "sales", "losses", "closing"):
        for r in form.get(key, []):
            a = agg.setdefault(r["item"], {})
            a[f"{key}_qty"] = a.get(f"{key}_qty", 0.0) + float(r.get("qty") or 0)
            a[f"{key}_val"] = a.get(f"{key}_val", 0.0) + float(r.get("value") or 0)
    rows = []
    for it in sorted(set(items) | set(agg)):
        a = agg.get(it, {})
        g = lambda k: a.get(k, 0.0)
        exp_q = g("opening_qty") + g("in_qty") - g("sales_qty") - g("losses_qty")
        exp_v = g("opening_val") + g("in_val") - g("sales_val") - g("losses_val")
        cat = (items.get(it) or {}).get("category", "")
        rows.append({"Item": it, "Name": (items.get(it) or {}).get("name", ""), "Category": cat,
                     "Category name": names.get(cat, "?"), "Unit": (items.get(it) or {}).get("unit", ""),
                     "Opening qty": g("opening_qty"), "In qty": g("in_qty"), "Sales qty": g("sales_qty"),
                     "Losses qty": g("losses_qty"), "Expected closing qty": exp_q, "Counted qty": g("closing_qty"),
                     "Qty difference": g("closing_qty") - exp_q,
                     "Opening value": g("opening_val"), "In value": g("in_val"), "Sales value (cost)": g("sales_val"),
                     "Losses value": g("losses_val"), "Expected closing value": exp_v, "Counted value": g("closing_val"),
                     "Value difference": g("closing_val") - exp_v})
    return pd.DataFrame(rows)


def by_category(mv: pd.DataFrame, chart: dict) -> pd.DataFrame:
    vals = ["Opening value", "In value", "Sales value (cost)", "Losses value", "Expected closing value", "Counted value",
            "Value difference"]
    rows = []
    for c in cats(chart):
        sub = mv[mv["Category"] == c["code"]] if not mv.empty else mv
        rows.append({"Code": c["code"], "Category": c["name"], **{v: float(sub[v].sum()) if len(sub) else 0.0 for v in vals}})
    df = pd.DataFrame(rows)
    df = pd.concat([df, pd.DataFrame([{"Code": "", "Category": "TOTAL", **df[vals].sum().to_dict()}])], ignore_index=True)
    return df


def _net(tb, codes, y):
    if tb is None or f"debit_{y}" not in tb.columns:
        return None
    sel = tb[tb["code"].isin(codes)]
    return float(sel[f"debit_{y}"].sum() - sel[f"credit_{y}"].sum())


def reconcile(form: dict, tb: pd.DataFrame, chart: dict, year: int) -> pd.DataFrame:
    mv = movement(form, chart)
    cat = by_category(mv, chart).set_index("Code")
    rows = []
    pl_codes = [x for c in cats(chart) for x in (c.get("pl_open"), c.get("pl_close")) if x]
    periodic = tb is not None and f"debit_{year}" in tb.columns and \
        float(tb[tb["code"].isin(pl_codes)][[f"debit_{year}", f"credit_{year}"]].abs().sum().sum()) > TOL
    for c in cats(chart):
        close_f, open_f = cat.at[c["code"], "Counted value"], cat.at[c["code"], "Opening value"]
        tb_close = _net(tb, [c["code"]], year)
        tb_open = _net(tb, [c["code"]], year - 1)
        if abs(close_f) > TOL or abs(tb_close or 0) > TOL:
            rows.append({"Item": f"Closing stock — {c['name']}", "Code": c["code"], "Stock form": close_f,
                         "Trial balance": tb_close})
        if abs(open_f) > TOL or abs(tb_open or 0) > TOL:
            rows.append({"Item": f"Opening stock — {c['name']} (TB of {year - 1})", "Code": c["code"], "Stock form": open_f,
                         "Trial balance": tb_open})
        if periodic and c.get("pl_close") and (abs(close_f) > TOL or abs(_net(tb, [c["pl_close"]], year) or 0) > TOL):
            rows.append({"Item": f"Closing stock in the P&L — {c['name']}", "Code": c["pl_close"], "Stock form": close_f,
                         "Trial balance": -(_net(tb, [c["pl_close"]], year) or 0)})
        if periodic and c.get("pl_open") and (abs(open_f) > TOL or abs(_net(tb, [c["pl_open"]], year) or 0) > TOL):
            rows.append({"Item": f"Opening stock in the P&L — {c['name']}", "Code": c["pl_open"], "Stock form": open_f,
                         "Trial balance": _net(tb, [c["pl_open"]], year)})
    purch_codes = ["PL 02.02.01", "PL 02.02.01.01"]
    rows.append({"Item": "Purchases of the year", "Code": " + ".join(purch_codes),
                 "Stock form": float(mv["In value"].sum()) if len(mv) else 0.0, "Trial balance": _net(tb, purch_codes, year)})
    df = pd.DataFrame(rows)
    df["Difference"] = df.apply(lambda r: None if r["Trial balance"] is None else r["Stock form"] - r["Trial balance"], axis=1)
    df["OK"] = df["Difference"].apply(lambda d: d is not None and abs(d) <= TOL)
    return df


def inventory_movement(form: dict, chart: dict) -> dict:
    """Figures for note 12 (inventory movement): the count difference and the losses are the gain/(loss)."""
    mv = movement(form, chart)
    s = lambda c: float(mv[c].sum()) if len(mv) else 0.0
    opening, purchases, sales, counted = s("Opening value"), s("In value"), s("Sales value (cost)"), s("Counted value")
    return {"opening": opening, "purchases": purchases, "purchase_returns": 0.0, "cost_of_sales": -sales,
            "gain_loss": counted - opening - purchases + sales}


def note12_from_project(project: dict, tb: pd.DataFrame, chart: dict, year: int) -> tuple[dict | None, str]:
    """Inventory movement of a year from the stock form, when its closing stock equals the TB inventories."""
    form = (project.get("stock") or {}).get(str(year))
    if not form or not form.get("items"):
        return None, "no stock form"
    rec = reconcile(form, tb, chart, year)
    close = rec[rec["Item"].str.startswith("Closing stock —")]
    if len(close) and not close["OK"].all():
        return None, "stock form NOT used for note 12: its closing stock differs from the TB inventories (9 · Stock Control)"
    return inventory_movement(form, chart), "note 12 from the stock form"


# -------------------------------------------------------------------- Excel --
def export_workbook(form: dict, chart: dict, year: int, tb: pd.DataFrame | None, company: str = "") -> bytes:
    """Form sheets (inputs) + movement by item and by category with SUMIF formulas + reconciliation with the TB."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    wb = load_workbook(io.BytesIO(template_workbook(chart, year, form)))
    h, hf = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    red = Font(name="Arial", bold=True, color="C00000")
    mv = wb.create_sheet("Movement by item", 1)
    hdr = ["Item code", "Name", "Category", "Opening value", "In value", "Sales value (cost)", "Losses value",
           "Expected closing value", "Counted value", "Difference (counted − expected)", "Opening qty", "In qty", "Sales qty",
           "Losses qty", "Expected qty", "Counted qty", "Qty difference"]
    mv.append(hdr)
    for c in mv[1]:
        c.font, c.fill = h, hf
    rng = {"Opening stock": ("A", "B", "D"), "IN - Purchases": ("B", "C", "E"), "OUT - Sales": ("B", "C", "E"),
           "OUT - Losses & damages": ("B", "C", "E"), "Closing stock": ("A", "B", "D")}

    def sumif(sheet, kind, r):
        item_col, qty_col, val_col = rng[sheet]
        col = val_col if kind == "v" else qty_col
        return f"SUMIF('{sheet}'!${item_col}:${item_col},$A{r},'{sheet}'!${col}:${col})"
    for i, it in enumerate(form.get("items", []), start=2):
        r = i
        mv.append([it["item"], it.get("name", ""), it.get("category", ""),
                   f"={sumif('Opening stock', 'v', r)}", f"={sumif('IN - Purchases', 'v', r)}",
                   f"={sumif('OUT - Sales', 'v', r)}", f"={sumif('OUT - Losses & damages', 'v', r)}",
                   f"=D{r}+E{r}-F{r}-G{r}", f"={sumif('Closing stock', 'v', r)}", f"=I{r}-H{r}",
                   f"={sumif('Opening stock', 'q', r)}", f"={sumif('IN - Purchases', 'q', r)}",
                   f"={sumif('OUT - Sales', 'q', r)}", f"={sumif('OUT - Losses & damages', 'q', r)}",
                   f"=K{r}+L{r}-M{r}-N{r}", f"={sumif('Closing stock', 'q', r)}", f"=P{r}-O{r}"])
        for cc in mv[r][3:]:
            cc.number_format = MONEY
    last = mv.max_row
    for j, w in enumerate((14, 30, 16) + (15,) * 14, start=1):
        mv.column_dimensions[chr(64 + j)].width = w
    mv.freeze_panes = "D2"
    ct = wb.create_sheet("By category", 2)
    ct.append(["CIT code", "Category", "Opening value", "In value", "Sales value (cost)", "Losses value",
               "Expected closing value", "Counted value", "Difference"])
    for c in ct[1]:
        c.font, c.fill = h, hf
    for c in cats(chart):
        r = ct.max_row + 1
        ct.append([c["code"], c["name"]] + [f"=SUMIF('Movement by item'!$C$2:$C${last},$A{r},'Movement by item'!{col}$2:{col}${last})"
                                            for col in "DEFGHIJ"])
        for cc in ct[r][2:]:
            cc.number_format = MONEY
    r = ct.max_row + 1
    ct.append(["", "TOTAL"] + [f"=SUM({col}2:{col}{r - 1})" for col in "CDEFGHI"])
    for cc in ct[r]:
        cc.font = Font(name="Arial", bold=True)
        cc.number_format = MONEY
    ct.column_dimensions["B"].width = 40
    for col in "CDEFGHI":
        ct.column_dimensions[col].width = 16
    if tb is not None:
        rec = reconcile(form, tb, chart, year)
        rs = wb.create_sheet("Reconciliation TB", 3)
        rs.append(list(rec.columns))
        for c in rs[1]:
            c.font, c.fill = h, hf
        for rw in rec.itertuples(index=False):
            rs.append(["" if v is None else v for v in rw])
            for cc in rs[rs.max_row][2:5]:
                cc.number_format = MONEY
            if not rw.OK:
                for cc in rs[rs.max_row]:
                    cc.font = red
        rs.column_dimensions["A"].width = 52
        for col in "BCDE":
            rs.column_dimensions[col].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
