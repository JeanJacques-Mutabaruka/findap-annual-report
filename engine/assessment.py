"""Global assessment of the annual-report production: status per area, key figures, missing elements, key events,
all controls — and a date-and-time stamped Word / PDF report (summary or detailed), for the audit file and the client.
"""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd

from engine import fixed_assets as fa
from engine import multiyear, rra, stock, tax_cit, tb_io

ICON = {"ok": "✅", "warn": "⚠️", "fail": "❌", "todo": "⏳", "na": "—"}
WORD = {"ok": "OK", "warn": "To check", "fail": "Blocking", "todo": "Missing", "na": "Not applicable"}


def build(project: dict, tb: pd.DataFrame, chart: dict, models: dict[int, dict], valid_codes: set[str],
          import_notes: list[str] | None = None, files_generated: bool = False, version: str = "") -> dict:
    """Everything the assessment page and report show."""
    now = datetime.now()
    meta = project.get("meta", {})
    ys = tb_io.years_of(tb) if tb is not None and not tb.empty else []
    areas, missing, events = [], [], []

    def area(name, status, detail):
        areas.append({"Area": name, "Status": status, "Detail": detail})

    # --- trial balance
    if tb is None or tb.empty:
        area("Trial balance", "todo", "No trial balance loaded")
        return {"now": now, "meta": meta, "areas": areas, "missing": ["Trial balance"], "events": [], "figures": pd.DataFrame(),
                "controls": pd.DataFrame(), "details": {}, "version": version, "years": []}
    bal = tb_io.balance(tb)
    unbal = [y for y in ys if not bal[y]["ok"]]
    area("Trial balance", "fail" if unbal else "ok",
         f"{len(tb)} lines, years {', '.join(map(str, ys))}" +
         (" — NOT balanced in " + ", ".join(f"{y} ({bal[y]['diff']:,.0f})" for y in unbal) if unbal else " — balanced every year"))
    if unbal:
        missing.append("A balanced trial balance for " + ", ".join(map(str, unbal)))
    notes = tb_io.tb_notices(tb)
    neg = [n for n in notes if n.startswith("NEGATIVE")]
    if neg:
        events.append(neg[0].replace("NEGATIVE AMOUNTS — ", "Negative amounts: "))
    for n in import_notes or []:
        if "V11 numbering" in n:
            events.append(n)
    bad = tb[~tb["code"].isin(valid_codes)]
    area("CIT codes", "fail" if len(bad) else "ok", "every line has a valid CIT code" if bad.empty else
         f"{len(bad)} line(s) without a valid CIT code")
    if len(bad):
        missing.append(f"CIT codes for {len(bad)} TB line(s)")
    iss = tb_io.line_issues(tb, valid_codes)
    nb = int((iss["Level"] == "BLOCKING").sum()) if not iss.empty else 0
    nw = int((iss["Level"] == "WARNING").sum()) if not iss.empty else 0
    area("TB line checks", "fail" if nb else ("warn" if nw else "ok"), f"{nb} blocking, {nw} warning(s)")
    # --- company data
    req = {"company_name": meta.get("company_name"), "period_end": meta.get("period_end")}
    g = project.get("general", {})
    miss_g = [k for k, v in {**req, "audit_company_name": g.get("audit_company_name"), "directors": g.get("directors"),
                             "registered_office": g.get("registered_office"), "signature_date": g.get("signature_date")}.items()
              if not str(v or "").strip()]
    area("Company & report data", "warn" if miss_g else "ok", "complete" if not miss_g else "missing: " + ", ".join(miss_g))
    if miss_g:
        missing.append("Company / report data: " + ", ".join(miss_g))
    # --- statements
    ctrl_rows, figs = [], []
    for y, m in sorted(models.items()):
        for c in m["controls"]:
            ctrl_rows.append({"Year": y, "Level": c["level"], "Control": c["id"], "Message": c["message"]})
        kf, tax = m["key_figures"], m["income_tax"]
        figs.append({"Year": y, "Revenue": kf["revenue"]["cy"], "Gross profit": kf["gross_profit"]["cy"],
                     "Profit before tax": kf["pbt"]["cy"], "Income tax": tax["charge"]["cy"],
                     "Net profit": kf["net_profit"]["cy"], "Total assets": kf["total_assets"]["cy"],
                     "Total equity": kf["total_equity"]["cy"], "Tax payable": tax["payable"]["cy"]})
    controls = pd.DataFrame(ctrl_rows, columns=["Year", "Level", "Control", "Message"])
    if models:
        nbc = int((controls["Level"] == "BLOCKING").sum())
        nwc = int((controls["Level"] == "WARNING").sum())
        area("Statements & controls", "fail" if nbc else ("warn" if nwc else "ok"),
             f"{len(models)} report(s) ({', '.join(map(str, sorted(models)))}) — {nbc} blocking, {nwc} warning(s)")
    else:
        area("Statements & controls", "todo", "statements not computed (company name and year end needed)")
    # --- tax
    details = {}
    last = max(models) if models else None
    if last:
        m = models[last]
        c = m["income_tax"]["detail"]["cy"]
        tw = tax_cit.tax_warnings(m)
        area("Taxable income", "warn" if len(tw) > 1 else "ok",
             f"{last}: tax base {c['base']:,.0f}, tax {c['charge']:,.0f}, payable {c['payable']:,.0f}")
        details["Taxable income " + str(last)] = tax_cit.taxable_income_df(m).drop(columns=["kind"])
        if c["lb"] < 0:
            events.append(f"{last}: tax losses of previous years deducted {-c['lb']:,.0f} — RRA line 15 is higher than the "
                          "tax base by this amount (RRA 15 does not deduct line 14).")
        if c.get("mgmt") and c["mgmt"]["amount"] > 0:
            events.append(f"{last}: management fees above {c['mgmt']['limit_rate']:.0%} of turnover added back: "
                          f"{c['mgmt']['amount']:,.0f}.")
        for d in c["deductions"]:
            if d["amount"] > 0:
                events.append(f"{last}: {d['label']} deducted from the taxable income: {d['amount']:,.0f}.")
        if not c["prepay"] and c["charge"] > 0:
            missing.append(f"Quarterly CIT prepayments {last} (if any)")
        # RRA annex
        rows = multiyear.pair_input(project, tb, last, chart)["trial_balance"]
        table = tax_cit.rra_df(chart, rra.load(), rows, m)
        rec = tax_cit.reconciliation_df(table, m)
        badr = rec[rec["Difference"].abs() > 1]
        area("RRA CIT annex", "fail" if len(badr) else "ok",
             "same key figures as the annual report" if badr.empty else f"{len(badr)} difference(s) with the report")
        details["RRA annex vs report"] = rec
    # --- fixed assets
    reg = fa.register_from_records(project.get("asset_register"))
    ppe_tb = sum(float(tb[tb["code"] == c_["code"]][f"debit_{max(ys)}"].sum()) for c_ in chart["fixed_assets"]["categories"])
    if reg.empty:
        area("Fixed assets", "warn" if abs(ppe_tb) > 1 else "na",
             "no asset register — note 11 allocates the TB depreciation pro rata to cost" if abs(ppe_tb) > 1 else
             "no fixed assets in the TB")
        if abs(ppe_tb) > 1:
            missing.append("Fixed-asset register (8 · Long-term Assets Check → template)")
    else:
        s = fa.settings_of(project)
        book = fa.book_schedule(reg, chart, s, max(ys))
        tax_s = fa.tax_schedule(reg, chart, s, max(ys))
        recf = fa.reconcile(book, tb, chart, ys)
        bad_y = [y for y in ys if not fa.year_ok(recf, y)]
        fiss = fa.issues(reg, chart, s)
        nbf = int((fiss["Level"] == "BLOCKING").sum()) if not fiss.empty else 0
        area("Fixed assets", "fail" if nbf else ("warn" if bad_y else "ok"),
             f"{len(reg)} asset(s) in the register — " + ("reconciled with the TB every year" if not bad_y else
                                                           "differences with the TB in " + ", ".join(map(str, bad_y))) +
             (f"; {nbf} asset(s) with blocking issues" if nbf else ""))
        details["Fixed assets — reconciliation with the TB"] = recf
        bt = fa.book_vs_tax(book, tax_s, ys)
        details["Fixed assets — book vs tax depreciation"] = bt
        for _, r in bt.iterrows():
            if abs(r["Difference (tax − book)"]) > 1:
                events.append(f"{int(r['Year'])}: tax depreciation (RRA pools) {r['Tax depreciation (RRA pools)']:,.0f} ≠ book "
                              f"depreciation {r['Book depreciation (straight-line)']:,.0f} — difference "
                              f"{r['Difference (tax − book)']:,.0f} not yet reflected in the tax computation.")
        disp = book[book["disposals"] != 0]
        for y, grp in disp.groupby("year"):
            events.append(f"{y}: {len(grp)} asset disposal(s), gain/(loss) {grp['gain_loss'].sum():,.0f}.")
        ts = fa.tax_summary(tax_s, chart)
        for _, r in ts[ts["written_off"]].iterrows() if not ts.empty else []:
            if r["year"] in ys:
                events.append(f"{r['year']}: tax pool '{r['name']}' below the limit — written off at 100% "
                              f"({r['allowance']:,.0f}).")
        mr = [c_ for c_ in fa.categories(chart, s) if c_["rate"] is None and (reg["category"] == c_["code"]).any()]
        if mr:
            missing.append("Depreciation rate for " + ", ".join(c_["name"] for c_ in mr))
    # --- stock
    forms = project.get("stock") or {}
    st_rows = []
    for y in (sorted(models) or ys):
        inv_tb = sum(float(tb[tb["code"] == c_["code"]][f"debit_{y}"].sum() - tb[tb["code"] == c_["code"]][f"credit_{y}"].sum())
                     for c_ in stock.cats(chart))
        form = forms.get(str(y))
        if form and form.get("items"):
            rec_s = stock.reconcile(form, tb, chart, y)
            mv = stock.movement(form, chart)
            loss = float(mv["Losses value"].sum())
            cdiff = float(mv["Value difference"].sum())
            st_rows.append((y, "ok" if rec_s["OK"].all() else "warn",
                            f"{y}: form reconciled" if rec_s["OK"].all() else f"{y}: {int((~rec_s['OK']).sum())} difference(s) with the TB"))
            details[f"Stock {y} — reconciliation with the TB"] = rec_s
            if loss:
                events.append(f"{y}: stock losses / damages declared {loss:,.0f} ({len(form.get('losses', []))} event(s)).")
            if abs(cdiff) > 1:
                events.append(f"{y}: stock count differs from the expected stock by {cdiff:,.0f} (unexplained).")
        elif abs(inv_tb) > 1:
            st_rows.append((y, "warn", f"{y}: inventories {inv_tb:,.0f} in the TB, no stock form"))
            missing.append(f"Stock form {y} (inventories of {inv_tb:,.0f} not explained)")
    if st_rows:
        worst = "warn" if any(s_ == "warn" for _, s_, _ in st_rows) else "ok"
        area("Stock", worst, "; ".join(d for _, _, d in st_rows))
    else:
        area("Stock", "na", "no inventories in the TB")
    # --- tax losses
    ls = multiyear.loss_schedule(project, tb, chart)
    if ls["origins"]:
        events.append("Tax losses by year of origin: " + ", ".join(f"{o['year']}: {o['loss']:,.0f} "
                                                                    f"(remaining {o['remaining']:,.0f})" for o in ls["origins"]))
    area("Files generated", "ok" if files_generated else "todo", "Word / Excel files produced" if files_generated else
         "not yet (5 · Generate & Download)")
    return {"now": now, "meta": meta, "areas": areas, "missing": missing, "events": events,
            "figures": pd.DataFrame(figs), "controls": controls, "details": details, "version": version, "years": ys,
            "line_issues": iss}


# ------------------------------------------------------------------ report ---
def to_docx(a: dict, detailed: bool = False) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(9.5)
    company = a["meta"].get("company_name") or "Company"
    t = doc.add_heading(f"{company} — Annual report production: {'detailed' if detailed else 'summary'} assessment", 0)
    for r in t.runs:
        r.font.size = Pt(18)
    p = doc.add_paragraph()
    p.add_run(f"Generated on {a['now']:%d %B %Y at %H:%M} · Annual Report Generator {a['version']} · "
              f"financial year end {a['meta'].get('period_end') or '—'} · years {', '.join(map(str, a['years']))}").italic = True

    def table(df: pd.DataFrame, money: list[str] | None = None, max_rows: int = 400):
        money = money or []
        df = df.head(max_rows)
        tb_ = doc.add_table(rows=1, cols=len(df.columns))
        tb_.style = "Light Grid Accent 1"
        for j, c in enumerate(df.columns):
            tb_.rows[0].cells[j].text = str(c)
        for _, row in df.iterrows():
            cells = tb_.add_row().cells
            for j, c in enumerate(df.columns):
                v = row[c]
                if c in money and isinstance(v, (int, float)) and not pd.isna(v):
                    txt = f"({abs(v):,.0f})" if v < -0.5 else ("-" if abs(v) < 0.5 else f"{v:,.0f}")
                    cells[j].text = txt
                    cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    if v < -0.5:
                        for run in cells[j].paragraphs[0].runs:
                            run.font.color.rgb = RGBColor(0xC0, 0, 0)
                else:
                    cells[j].text = "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)
        for row in tb_.rows:
            for c in row.cells:
                for pp in c.paragraphs:
                    for run in pp.runs:
                        run.font.size = Pt(8)
        doc.add_paragraph()

    doc.add_heading("1. Status by area", 1)
    ar = pd.DataFrame(a["areas"])
    ar["Status"] = ar["Status"].map(lambda s: WORD.get(s, s))
    table(ar)
    doc.add_heading("2. Key figures", 1)
    if not a["figures"].empty:
        f = a["figures"].copy()
        f["Year"] = f["Year"].astype(str)
        table(f, [c for c in f.columns if c != "Year"])
    doc.add_heading("3. Missing elements", 1)
    for m in a["missing"] or ["None."]:
        par = doc.add_paragraph(m, style="List Bullet")
        if a["missing"]:
            for run in par.runs:
                run.font.color.rgb = RGBColor(0xC0, 0, 0)
    doc.add_heading("4. Key events and points to highlight", 1)
    for e in a["events"] or ["None."]:
        doc.add_paragraph(e, style="List Bullet")
    ctl = a["controls"]
    doc.add_heading("5. Controls", 1)
    if ctl.empty:
        doc.add_paragraph("No controls (statements not computed).")
    else:
        sub = ctl if detailed else ctl[ctl["Level"].isin(["BLOCKING", "WARNING"])]
        doc.add_paragraph(f"{int((ctl['Level'] == 'BLOCKING').sum())} blocking, {int((ctl['Level'] == 'WARNING').sum())} "
                          f"warning(s), {int((ctl['Level'] == 'INFO').sum())} information(s)." +
                          ("" if detailed else " Blocking and warnings listed below."))
        table(sub)
    if detailed:
        doc.add_heading("6. Details", 1)
        for title, df in a["details"].items():
            doc.add_heading(title, 2)
            money = [c for c in df.columns if df[c].dtype.kind in "fi" and c not in ("Year", "Coefficient", "RRA row")]
            table(df, money)
        li = a.get("line_issues")
        if li is not None and not li.empty:
            doc.add_heading("Trial-balance line issues", 2)
            table(li, max_rows=300)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(docx_bytes: bytes) -> bytes | None:
    """PDF through LibreOffice (same converter as the annual report); None when LibreOffice is absent."""
    import tempfile
    from pathlib import Path
    from engine import export_pdf
    if not export_pdf.find_soffice():
        return None
    with tempfile.TemporaryDirectory() as d:
        src, out = Path(d) / "assessment.docx", Path(d) / "assessment.pdf"
        src.write_bytes(docx_bytes)
        export_pdf.export(src, out)
        return out.read_bytes() if out.exists() else None
